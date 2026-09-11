from __future__ import annotations

import copy
import io
import zipfile

import pytest

from personagraph.input_processing.documents.contracts import (
    DiagnosticCode,
    DocumentNonTextKind,
    DocumentPageInventoryStatus,
    DocumentPageState,
    ElementKind,
    ProcessingAdmissionStatus,
)
from personagraph.input_processing.documents.readers import (
    UnsupportedLegacyOfficeFormat,
    read_document,
    supported_suffixes,
)
from personagraph.input_processing.documents.readers import office
from personagraph.input_processing.documents.readers.office import (
    MAX_ELEMENT_CHARS,
    read_docx,
    read_pptx,
)


def _empty_zip_payload() -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w"):
        pass
    return output.getvalue()


@pytest.fixture
def heading_docx(tmp_path):
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_heading("方案", level=1)
    document.add_paragraph("总体思路是先做入口。")
    document.add_heading("风险", level=2)
    document.add_paragraph("扫描件无法直接读取。")
    path = tmp_path / "plan.docx"
    document.save(str(path))
    return path


def test_docx_keeps_its_heading_path_instead_of_a_paragraph_counter(heading_docx):
    result = read_docx(heading_docx)

    body = [e for e in result.elements if e.kind is ElementKind.PARAGRAPH]
    assert body[0].locator.section_path == ("方案",)
    assert body[1].locator.section_path == ("方案", "风险")
    # Word 在此层级没有页概念，因此明确说明，而不是臆造页码。
    assert all(element.locator.page is None for element in result.elements)
    assert "段" not in body[0].locator.describe()


def test_docx_keeps_paragraphs_and_tables_in_document_order(tmp_path):
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph("before-table")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "left"
    table.cell(0, 1).text = "right"
    document.add_paragraph("after-table")
    path = tmp_path / "ordered.docx"
    document.save(str(path))

    result = read_docx(path)

    assert [(element.kind, element.text) for element in result.elements] == [
        (ElementKind.PARAGRAPH, "before-table"),
        (ElementKind.TABLE, "left | right"),
        (ElementKind.PARAGRAPH, "after-table"),
    ]


def test_docx_long_blocks_are_losslessly_segmented(tmp_path):
    docx = pytest.importorskip("docx")
    expected = "A" * MAX_ELEMENT_CHARS + "TAIL-SENTINEL"
    document = docx.Document()
    document.add_paragraph(expected)
    path = tmp_path / "long.docx"
    document.save(str(path))

    result = read_docx(path)

    assert "".join(element.text or "" for element in result.elements) == expected
    assert all(len(element.text or "") <= MAX_ELEMENT_CHARS for element in result.elements)
    assert DiagnosticCode.LIMIT_REACHED not in {item.code for item in result.diagnostics}
    assert [element.locator.char_range for element in result.elements] == [
        (0, MAX_ELEMENT_CHARS),
        (MAX_ELEMENT_CHARS, len(expected)),
    ]


def test_docx_embedded_picture_is_a_located_visual_gap(tmp_path):
    docx = pytest.importorskip("docx")
    image_path = _tiny_png(tmp_path)
    document = docx.Document()
    document.add_heading("Evidence", level=1)
    document.add_paragraph("before-picture")
    document.add_picture(str(image_path))
    document.add_paragraph("after-picture")
    path = tmp_path / "mixed.docx"
    document.save(str(path))

    result = read_docx(path)

    assert [element.text for element in result.elements] == [
        "Evidence",
        "before-picture",
        None,
        "after-picture",
    ]
    visual = result.elements[2]
    assert visual.kind is ElementKind.IMAGE
    assert visual.needs_vision is True
    assert visual.locator.section_path == ("Evidence",)
    assert any(
        item.code is DiagnosticCode.PAGE_NEEDS_VISION
        and item.locator == visual.locator
        for item in result.diagnostics
    )
    assert result.admission_status is ProcessingAdmissionStatus.PARTIAL


def test_docx_omml_formula_is_never_silently_dropped(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    document = docx.Document()
    document.add_paragraph("BODY")
    formula = document.add_paragraph()
    formula._p.append(parse_xml(
        f'<m:oMathPara {nsdecls("m")}>'
        '<m:oMath><m:r><m:t>x²+y²=z²</m:t></m:r></m:oMath>'
        '</m:oMathPara>'
    ))
    path = tmp_path / "formula.docx"
    document.save(str(path))

    result = read_docx(path)

    assert [element.text for element in result.elements] == ["BODY", None]
    assert result.elements[-1].needs_vision is True
    assert any(
        item.code is DiagnosticCode.PAGE_NEEDS_VISION
        and "formula" in (item.detail or "").lower()
        for item in result.diagnostics
    )
    assert result.admission_status is ProcessingAdmissionStatus.PARTIAL


def test_docx_includes_distinct_header_and_footer_content_once(tmp_path):
    docx = pytest.importorskip("docx")
    document = docx.Document()
    section = document.sections[0]
    section.header.paragraphs[0].text = "CONFIDENTIAL HEADER"
    section.footer.paragraphs[0].text = "PAGE FOOTER"
    document.add_paragraph("BODY CONTENT")
    path = tmp_path / "header-footer.docx"
    document.save(str(path))

    result = read_docx(path)

    assert [element.text for element in result.elements] == [
        "CONFIDENTIAL HEADER",
        "BODY CONTENT",
        "PAGE FOOTER",
    ]
    assert result.elements[0].locator.section_path == ("header:default",)
    assert result.elements[-1].locator.section_path == ("footer:default",)
    assert result.admission_status is ProcessingAdmissionStatus.COMPLETE


def test_linked_docx_headers_are_not_duplicated_across_sections(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.enum.section import WD_SECTION

    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "SHARED HEADER"
    document.add_paragraph("FIRST BODY")
    second = document.add_section(WD_SECTION.NEW_PAGE)
    assert second.header.is_linked_to_previous is True
    document.add_paragraph("SECOND BODY")
    path = tmp_path / "linked-header.docx"
    document.save(str(path))

    result = read_docx(path)

    assert [element.text for element in result.elements].count("SHARED HEADER") == 1


def test_docx_omits_inactive_first_and_even_header_variants(tmp_path):
    docx = pytest.importorskip("docx")

    document = docx.Document()
    section = document.sections[0]
    section.different_first_page_header_footer = False
    document.settings.odd_and_even_pages_header_footer = False
    section.header.paragraphs[0].text = "ACTIVE DEFAULT HEADER"
    section.first_page_header.paragraphs[0].text = "HIDDEN FIRST HEADER 73500"
    section.even_page_header.paragraphs[0].text = "HIDDEN EVEN HEADER 84600"
    document.add_paragraph("VISIBLE BODY")
    path = tmp_path / "inactive-header-variants.docx"
    document.save(str(path))

    result = read_docx(path)
    text = "\n".join(element.text or "" for element in result.elements)

    assert "ACTIVE DEFAULT HEADER" in text
    assert "VISIBLE BODY" in text
    assert "HIDDEN FIRST HEADER 73500" not in text
    assert "HIDDEN EVEN HEADER 84600" not in text
    assert result.admission_status is ProcessingAdmissionStatus.COMPLETE


def test_docx_content_controls_are_unwrapped_in_body_order(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.oxml import OxmlElement

    document = docx.Document()
    document.add_paragraph("BEFORE")
    paragraph = document.add_paragraph("CONTROL-SENTINEL")
    body = document.element.body
    body.remove(paragraph._p)
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    content.append(paragraph._p)
    control.append(content)
    body.insert(len(body) - 1, control)
    document.add_paragraph("AFTER")
    path = tmp_path / "content-control.docx"
    document.save(str(path))

    result = read_docx(path)

    assert [element.text for element in result.elements] == [
        "BEFORE",
        "CONTROL-SENTINEL",
        "AFTER",
    ]


def test_docx_current_visible_tracked_insertions_are_included(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.oxml import OxmlElement

    document = docx.Document()
    document.add_paragraph("BODY")
    inserted = document.add_paragraph("INSERTED CRITICAL 73500")
    body = document.element.body
    body.remove(inserted._p)
    wrapper = OxmlElement("w:ins")
    wrapper.append(inserted._p)
    body.insert(len(body) - 1, wrapper)
    path = tmp_path / "tracked.docx"
    document.save(str(path))

    result = read_docx(path)

    assert [element.text for element in result.elements] == [
        "BODY",
        "INSERTED CRITICAL 73500",
    ]


@pytest.mark.parametrize("wrapper_name", ("w:ins", "w:moveTo", "w:sdt"))
def test_docx_current_visible_run_wrappers_are_included(tmp_path, wrapper_name):
    docx = pytest.importorskip("docx")
    from docx.oxml import OxmlElement

    document = docx.Document()
    paragraph = document.add_paragraph("BODY ")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "CRITICAL 73500"
    run.append(text)
    wrapper = OxmlElement(wrapper_name)
    if wrapper_name == "w:sdt":
        content = OxmlElement("w:sdtContent")
        content.append(run)
        wrapper.append(content)
    else:
        wrapper.append(run)
    paragraph._p.append(wrapper)
    path = tmp_path / f"run-wrapper-{wrapper_name.removeprefix('w:')}.docx"
    document.save(str(path))

    result = read_docx(path)

    assert [element.text for element in result.elements] == [
        "BODY CRITICAL 73500"
    ]
    assert result.admission_status is ProcessingAdmissionStatus.COMPLETE


def test_docx_altchunk_is_an_explicit_gap_instead_of_silent_complete(tmp_path):
    docx = pytest.importorskip("docx")
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    document = docx.Document()
    document.add_paragraph("BODY")
    alt_chunk = OxmlElement("w:altChunk")
    alt_chunk.set(qn("r:id"), "rIdAlternativeContent")
    document.element.body.insert(len(document.element.body) - 1, alt_chunk)
    path = tmp_path / "altchunk.docx"
    document.save(str(path))

    result = read_docx(path)

    assert [element.text for element in result.text_elements()] == ["BODY"]
    assert any(element.needs_vision for element in result.elements)
    assert any(
        diagnostic.code is DiagnosticCode.PAGE_NEEDS_VISION
        and "alternative-format" in str(diagnostic.detail)
        for diagnostic in result.diagnostics
    )
    assert result.admission_status is ProcessingAdmissionStatus.PARTIAL


def test_docx_nested_table_text_is_not_silently_dropped(tmp_path):
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph("BODY")
    outer = document.add_table(rows=1, cols=1)
    outer.cell(0, 0).paragraphs[0].text = "OUTER"
    nested = outer.cell(0, 0).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "NESTED CRITICAL 73500"
    path = tmp_path / "nested-table.docx"
    document.save(str(path))

    result = read_docx(path)
    text = "\n".join(element.text or "" for element in result.elements)

    assert "OUTER" in text
    assert "NESTED CRITICAL 73500" in text
    assert result.admission_status is ProcessingAdmissionStatus.COMPLETE


def test_docx_referenced_footnote_story_is_included(tmp_path):
    docx = pytest.importorskip("docx")
    import zipfile
    from lxml import etree

    base = tmp_path / "footnote-base.docx"
    document = docx.Document()
    document.add_paragraph("Body claim")
    document.save(str(base))
    with zipfile.ZipFile(base) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}

    word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    document_xml = etree.fromstring(members["word/document.xml"])
    paragraph = document_xml.find(f".//{{{word_ns}}}body/{{{word_ns}}}p")
    run = etree.SubElement(paragraph, f"{{{word_ns}}}r")
    etree.SubElement(run, f"{{{word_ns}}}footnoteReference").set(
        f"{{{word_ns}}}id", "1"
    )
    members["word/document.xml"] = etree.tostring(
        document_xml, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    relationships = etree.fromstring(members["word/_rels/document.xml.rels"])
    relation = etree.SubElement(relationships, f"{{{rel_ns}}}Relationship")
    relation.set("Id", "rIdFootnotes")
    relation.set(
        "Type",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes",
    )
    relation.set("Target", "footnotes.xml")
    members["word/_rels/document.xml.rels"] = etree.tostring(
        relationships, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    content_types = etree.fromstring(members["[Content_Types].xml"])
    override = etree.SubElement(content_types, f"{{{content_ns}}}Override")
    override.set("PartName", "/word/footnotes.xml")
    override.set(
        "ContentType",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
    )
    members["[Content_Types].xml"] = etree.tostring(
        content_types, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    members["word/footnotes.xml"] = (
        f'<w:footnotes xmlns:w="{word_ns}">'
        '<w:footnote w:id="1"><w:p><w:r><w:t>'
        'CRITICAL FOOTNOTE EVIDENCE 73500'
        '</w:t></w:r></w:p></w:footnote></w:footnotes>'
    ).encode("utf-8")

    path = tmp_path / "footnote.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)

    result = read_docx(path)

    assert "CRITICAL FOOTNOTE EVIDENCE 73500" in [
        element.text for element in result.elements
    ]
    footnote = next(
        element for element in result.elements
        if element.text == "CRITICAL FOOTNOTE EVIDENCE 73500"
    )
    assert footnote.locator.section_path == ("footnote:1",)
    assert result.admission_status is ProcessingAdmissionStatus.COMPLETE


def test_docx_element_budget_applies_to_tables_too(tmp_path, monkeypatch):
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph("first")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "must-not-be-silent"
    path = tmp_path / "budget.docx"
    document.save(str(path))
    monkeypatch.setattr(office, "MAX_ELEMENTS", 1)

    result = read_docx(path)

    assert [element.text for element in result.elements] == ["first"]
    assert DiagnosticCode.LIMIT_REACHED in {item.code for item in result.diagnostics}
    assert result.admission_status is ProcessingAdmissionStatus.REJECTED


@pytest.mark.parametrize(
    ("payload", "expected_code"),
    [
        (b"not-a-zip", DiagnosticCode.CORRUPT_SOURCE),
        (_empty_zip_payload(), DiagnosticCode.CORRUPT_SOURCE),
    ],
)
def test_docx_container_preflight_returns_a_typed_rejection(
    tmp_path,
    payload,
    expected_code,
):
    path = tmp_path / "unsafe.docx"
    path.write_bytes(payload)

    result = read_docx(path)

    assert result.elements == ()
    assert [item.code for item in result.diagnostics] == [expected_code]
    assert result.admission_status is ProcessingAdmissionStatus.REJECTED


@pytest.fixture
def deck(tmp_path):
    pptx = pytest.importorskip("pptx")
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "第一页标题"
    path = tmp_path / "deck.pptx"
    presentation.save(str(path))
    return path


def test_pptx_slides_are_pages_in_the_shared_vocabulary(deck):
    result = read_pptx(deck)

    assert result.elements[0].locator.page == 1
    assert result.elements[0].locator.describe().startswith("p1")
    assert "slide" not in result.elements[0].locator.describe()


def test_pptx_long_text_is_losslessly_segmented(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    expected = "B" * MAX_ELEMENT_CHARS + "TAIL-SENTINEL"
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = expected
    path = tmp_path / "long.pptx"
    presentation.save(str(path))

    result = read_pptx(path)

    assert "".join(element.text or "" for element in result.elements) == expected
    assert all(len(element.text or "") <= MAX_ELEMENT_CHARS for element in result.elements)
    assert DiagnosticCode.LIMIT_REACHED not in {item.code for item in result.diagnostics}
    assert [element.locator.char_range for element in result.elements] == [
        (0, MAX_ELEMENT_CHARS),
        (MAX_ELEMENT_CHARS, len(expected)),
    ]


def test_pptx_extracts_table_cells_as_a_table_element(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(
        2, 2, Inches(1), Inches(1), Inches(5), Inches(2)
    ).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "92%"
    path = tmp_path / "table.pptx"
    presentation.save(str(path))

    result = read_pptx(path)

    assert [(element.kind, element.text) for element in result.elements] == [
        (ElementKind.TABLE, "Metric | Value\nAccuracy | 92%")
    ]


def test_pptx_text_slide_reports_picture_chart_and_group_as_visual_gaps(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(0), Inches(0), Inches(4), Inches(0.5)).text = "Summary"
    slide.shapes.add_picture(str(_tiny_png(tmp_path)), Inches(0), Inches(1))
    chart_data = ChartData()
    chart_data.categories = ["A", "B"]
    chart_data.add_series("Series", (1, 2))
    slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(2),
        Inches(1),
        Inches(3),
        Inches(2),
        chart_data,
    )
    group = slide.shapes.add_group_shape()
    group.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(1), Inches(1)
    ).text = "Grouped label"
    path = tmp_path / "visuals.pptx"
    presentation.save(str(path))

    result = read_pptx(path)

    assert "Summary" in [element.text for element in result.elements]
    assert "Grouped label" in [element.text for element in result.elements]
    visual_gaps = [element for element in result.elements if element.needs_vision]
    assert len(visual_gaps) == 3
    assert all(element.kind is ElementKind.IMAGE for element in visual_gaps)
    assert all(element.locator.page == 1 for element in visual_gaps)
    visual_diagnostics = [
        item for item in result.diagnostics
        if item.code is DiagnosticCode.PAGE_NEEDS_VISION
    ]
    assert [item.locator for item in visual_diagnostics] == [
        element.locator for element in visual_gaps
    ]
    assert result.admission_status is ProcessingAdmissionStatus.PARTIAL


def test_pptx_manifest_inventories_every_physical_slide_and_visual_evidence(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.chart.data import ChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    presentation = pptx.Presentation()
    mixed = presentation.slides.add_slide(presentation.slide_layouts[6])
    mixed.shapes.add_textbox(Inches(0), Inches(0), Inches(4), Inches(0.5)).text = (
        "Evidence slide"
    )
    mixed.shapes.add_picture(str(_tiny_png(tmp_path)), Inches(0), Inches(1))
    chart_data = ChartData()
    chart_data.categories = ["A", "B"]
    chart_data.add_series("Series", (1, 2))
    mixed.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        Inches(2),
        Inches(1),
        Inches(3),
        Inches(2),
        chart_data,
    )
    presentation.slides.add_slide(presentation.slide_layouts[6])
    final = presentation.slides.add_slide(presentation.slide_layouts[6])
    final.shapes.add_textbox(Inches(0), Inches(0), Inches(4), Inches(0.5)).text = (
        "Final slide"
    )
    path = tmp_path / "manifest.pptx"
    presentation.save(str(path))

    result = read_pptx(path)
    replay = read_pptx(path)

    manifest = result.page_manifest
    assert manifest is not None
    assert manifest.physical_page_count == 3
    assert manifest.inventory_status is DocumentPageInventoryStatus.COMPLETE
    assert [page.state for page in manifest.pages] == [
        DocumentPageState.MIXED,
        DocumentPageState.NO_EXTRACTABLE_CONTENT,
        DocumentPageState.TEXT,
    ]
    visual_elements = tuple(element for element in result.elements if element.needs_vision)
    units = manifest.pages[0].nontext_units
    elements_by_id = {element.element_id: element for element in result.elements}
    assert len(visual_elements) == len(units) == 2
    assert {unit.kind for unit in units} == {DocumentNonTextKind.FIGURE}
    assert {unit.element_id for unit in units} == {
        element.element_id for element in visual_elements
    }
    assert all(unit.locator == elements_by_id[unit.element_id].locator for unit in units)
    assert manifest.pages[1].diagnostics == (
        next(
            diagnostic
            for diagnostic in result.diagnostics
            if diagnostic.code is DiagnosticCode.PAGE_EMPTY
        ),
    )
    assert replay.page_manifest == manifest


def test_pptx_includes_effective_master_and_layout_nonplaceholder_shapes(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    presentation = pptx.Presentation()
    layout = presentation.slide_layouts[6]
    slide = presentation.slides.add_slide(layout)

    master_text = slide.shapes.add_textbox(
        Inches(0), Inches(0), Inches(4), Inches(0.4)
    )
    master_text.text = "MASTER LEGAL NOTICE"
    _move_shape_to(master_text, layout.slide_master.shapes)

    layout_logo = slide.shapes.add_shape(
        MSO_SHAPE.HEXAGON, Inches(6), Inches(0), Inches(1), Inches(1)
    )
    _move_shape_to(layout_logo, layout.shapes)

    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)
    ).text = "SLIDE BODY"
    path = tmp_path / "inherited-shapes.pptx"
    presentation.save(str(path))

    result = read_pptx(path)

    text_by_path = {
        element.locator.section_path: element.text
        for element in result.text_elements()
    }
    assert text_by_path[("slide_master",)] == "MASTER LEGAL NOTICE"
    assert "SLIDE BODY" in text_by_path.values()
    inherited_visuals = [
        element
        for element in result.elements
        if element.needs_vision
        and element.locator.section_path == ("slide_layout",)
    ]
    assert len(inherited_visuals) == 1
    assert result.page_manifest.pages[0].state is DocumentPageState.MIXED
    assert inherited_visuals[0].element_id in {
        unit.element_id for unit in result.page_manifest.pages[0].nontext_units
    }


def test_pptx_master_graphics_flag_suppresses_inherited_shapes(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    layout = presentation.slide_layouts[6]
    slide = presentation.slides.add_slide(layout)
    inherited = slide.shapes.add_textbox(
        Inches(0), Inches(0), Inches(4), Inches(0.4)
    )
    inherited.text = "MUST BE HIDDEN"
    _move_shape_to(inherited, layout.shapes)
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)
    ).text = "VISIBLE BODY"
    slide._element.set("showMasterSp", "0")
    path = tmp_path / "hidden-master-graphics.pptx"
    presentation.save(str(path))

    result = read_pptx(path)

    assert [element.text for element in result.text_elements()] == ["VISIBLE BODY"]


def test_pptx_does_not_emit_inherited_template_placeholder_prompts(tmp_path):
    pptx = pytest.importorskip("pptx")

    presentation = pptx.Presentation()
    layout = presentation.slide_layouts[0]
    layout.placeholders[0].text = "CLICK TO EDIT MASTER TITLE"
    slide = presentation.slides.add_slide(layout)
    slide.shapes.title.text = "ACTUAL TITLE"
    path = tmp_path / "placeholder-filter.pptx"
    presentation.save(str(path))

    result = read_pptx(path)

    texts = [element.text for element in result.text_elements()]
    assert "ACTUAL TITLE" in texts
    assert "CLICK TO EDIT MASTER TITLE" not in texts


def test_pptx_picture_placeholder_is_still_a_visual_gap(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.enum.shapes import PP_PLACEHOLDER

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[8])
    slide.shapes.title.text = "VISIBLE TITLE"
    picture_placeholder = next(
        shape for shape in slide.placeholders
        if shape.placeholder_format.type is PP_PLACEHOLDER.PICTURE
    )
    picture_placeholder.insert_picture(str(_tiny_png(tmp_path)))
    path = tmp_path / "picture-placeholder.pptx"
    presentation.save(str(path))

    result = read_pptx(path)

    assert "VISIBLE TITLE" in [element.text for element in result.elements]
    assert len([element for element in result.elements if element.needs_vision]) == 1
    assert result.admission_status is ProcessingAdmissionStatus.PARTIAL


def test_pptx_autoshape_keeps_text_and_visual_relationship_gap(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    arrow = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_ARROW,
        Inches(1),
        Inches(1),
        Inches(4),
        Inches(1),
    )
    arrow.text = "CAUSE TO EFFECT"
    path = tmp_path / "arrow.pptx"
    presentation.save(str(path))

    result = read_pptx(path)

    assert "CAUSE TO EFFECT" in [element.text for element in result.elements]
    assert len([element for element in result.elements if element.needs_vision]) == 1
    page = result.page_manifest.pages[0]
    assert page.nontext_units[0].text_element_ids == (page.text_element_ids[0],)
    assert result.admission_status is ProcessingAdmissionStatus.PARTIAL


def test_pptx_image_background_is_an_explicit_visual_gap(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches
    from pptx.oxml.ns import qn

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "VISIBLE TITLE"
    picture = slide.shapes.add_picture(
        str(_tiny_png(tmp_path)),
        Inches(1),
        Inches(1),
    )
    relationship_id = picture._element.blipFill.blip.rEmbed
    slide.shapes._spTree.remove(picture._element)
    background = OxmlElement("p:bg")
    background_properties = OxmlElement("p:bgPr")
    blip_fill = OxmlElement("a:blipFill")
    blip = OxmlElement("a:blip")
    blip.set(qn("r:embed"), relationship_id)
    stretch = OxmlElement("a:stretch")
    stretch.append(OxmlElement("a:fillRect"))
    blip_fill.append(blip)
    blip_fill.append(stretch)
    background_properties.append(blip_fill)
    background.append(background_properties)
    slide._element.cSld.insert(0, background)
    path = tmp_path / "image-background.pptx"
    presentation.save(str(path))

    result = read_pptx(path)

    assert "VISIBLE TITLE" in [element.text for element in result.text_elements()]
    visual = next(element for element in result.elements if element.needs_vision)
    assert visual.locator.section_path == ("slide_background",)
    assert result.page_manifest.pages[0].nontext_units[0].kind is (
        DocumentNonTextKind.FIGURE
    )
    assert result.admission_status is ProcessingAdmissionStatus.PARTIAL


def test_pptx_includes_speaker_notes_with_their_slide_locator(tmp_path):
    pptx = pytest.importorskip("pptx")
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "VISIBLE SLIDE"
    slide.notes_slide.notes_text_frame.text = "SPEAKER NOTE EVIDENCE"
    path = tmp_path / "notes.pptx"
    presentation.save(str(path))

    result = read_pptx(path)

    assert [element.text for element in result.elements] == [
        "VISIBLE SLIDE",
        "SPEAKER NOTE EVIDENCE",
    ]
    note = result.elements[-1]
    assert note.locator.page == 1
    assert note.locator.section_path == ("speaker_notes",)
    assert result.admission_status is ProcessingAdmissionStatus.COMPLETE


def test_pptx_includes_additional_notes_textboxes(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "VISIBLE SLIDE"
    slide.notes_slide.notes_text_frame.text = "ORDINARY NOTES"
    extra = slide.shapes.add_textbox(
        Inches(1),
        Inches(1),
        Inches(4),
        Inches(1),
    )
    extra.text = "EXTRA NOTES TEXT 73500"
    _move_shape_to(extra, slide.notes_slide.shapes)
    path = tmp_path / "extra-notes.pptx"
    presentation.save(str(path))

    result = read_pptx(path)

    assert "EXTRA NOTES TEXT 73500" in [
        element.text for element in result.text_elements()
    ]
    extra_note = next(
        element
        for element in result.text_elements()
        if element.text == "EXTRA NOTES TEXT 73500"
    )
    assert extra_note.locator.section_path == ("speaker_notes_extra",)
    assert result.admission_status is ProcessingAdmissionStatus.COMPLETE


def test_pptx_reads_extra_note_shapes_when_notes_body_placeholder_is_absent(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "VISIBLE SLIDE"
    notes_slide = slide.notes_slide
    for shape in tuple(notes_slide.shapes):
        if (
            getattr(shape, "is_placeholder", False)
            and shape.placeholder_format.type is PP_PLACEHOLDER.BODY
        ):
            shape._element.getparent().remove(shape._element)
    extra = slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)
    )
    extra.text = "EXTRA NOTE 73500"
    _move_shape_to(extra, notes_slide.shapes)
    path = tmp_path / "notes-without-body.pptx"
    presentation.save(str(path))

    result = read_pptx(path)

    assert [element.text for element in result.text_elements()] == [
        "VISIBLE SLIDE",
        "EXTRA NOTE 73500",
    ]
    assert result.admission_status is ProcessingAdmissionStatus.COMPLETE


def test_pptx_includes_classic_review_comment_text(tmp_path):
    pptx = pytest.importorskip("pptx")
    from lxml import etree

    base = tmp_path / "comment-base.pptx"
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "VISIBLE SLIDE"
    presentation.save(str(base))
    with zipfile.ZipFile(base) as archive:
        members = {name: archive.read(name) for name in archive.namelist()}

    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    rel_path = "ppt/slides/_rels/slide1.xml.rels"
    relationships = etree.fromstring(members[rel_path])
    relation = etree.SubElement(relationships, f"{{{rel_ns}}}Relationship")
    relation.set("Id", "rIdReviewComments")
    relation.set(
        "Type",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
    )
    relation.set("Target", "../comments/comment1.xml")
    members[rel_path] = etree.tostring(
        relationships, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    content_types = etree.fromstring(members["[Content_Types].xml"])
    override = etree.SubElement(content_types, f"{{{content_ns}}}Override")
    override.set("PartName", "/ppt/comments/comment1.xml")
    override.set(
        "ContentType",
        "application/vnd.openxmlformats-officedocument.presentationml.comments+xml",
    )
    members["[Content_Types].xml"] = etree.tostring(
        content_types, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    members["ppt/comments/comment1.xml"] = (
        '<p:cmLst xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main">'
        '<p:cm authorId="0" dt="2026-08-20T00:00:00Z" idx="1">'
        '<p:pos x="0" y="0"/><p:text>REVIEW COMMENT 73500</p:text>'
        '</p:cm></p:cmLst>'
    ).encode("utf-8")

    path = tmp_path / "comment.pptx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)

    result = read_pptx(path)

    comment = next(
        element
        for element in result.text_elements()
        if element.text == "REVIEW COMMENT 73500"
    )
    assert comment.locator.page == 1
    assert comment.locator.section_path[0] == "review_comment"
    assert result.admission_status is ProcessingAdmissionStatus.COMPLETE


def test_empty_pptx_slide_is_a_typed_gap_instead_of_silent_complete(tmp_path):
    pptx = pytest.importorskip("pptx")
    presentation = pptx.Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    path = tmp_path / "empty-slide.pptx"
    presentation.save(str(path))

    result = read_pptx(path)

    assert result.elements == ()
    assert any(
        diagnostic.code is DiagnosticCode.PAGE_EMPTY
        and diagnostic.locator.page == 1
        for diagnostic in result.diagnostics
    )
    assert result.admission_status is ProcessingAdmissionStatus.REJECTED


def test_pptx_element_budget_is_global_across_slides(tmp_path, monkeypatch):
    pptx = pytest.importorskip("pptx")
    presentation = pptx.Presentation()
    for text in ("first", "must-not-be-silent"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = text
    path = tmp_path / "budget.pptx"
    presentation.save(str(path))
    monkeypatch.setattr(office, "MAX_ELEMENTS", 1)

    result = read_pptx(path)

    assert [element.text for element in result.elements] == ["first"]
    assert DiagnosticCode.LIMIT_REACHED in {item.code for item in result.diagnostics}
    assert result.page_manifest.physical_page_count == 2
    assert result.page_manifest.inventory_status is DocumentPageInventoryStatus.PARTIAL
    assert result.page_manifest.pages[1].state is DocumentPageState.UNREADABLE
    assert result.page_manifest.pages[1].diagnostics[0].code is DiagnosticCode.LIMIT_REACHED
    assert result.admission_status is ProcessingAdmissionStatus.REJECTED


def test_same_name_and_size_decks_with_different_pixels_get_different_visual_ids(
    tmp_path,
):
    pptx = pytest.importorskip("pptx")
    image = pytest.importorskip("PIL.Image")
    from pptx.util import Inches

    encoded_decks: list[bytes] = []
    for color in ((255, 0, 0), (0, 0, 255)):
        picture = tmp_path / "pixel.bmp"
        image.new("RGB", (32, 32), color=color).save(picture)
        presentation = pptx.Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_picture(str(picture), Inches(1), Inches(1))
        output = io.BytesIO()
        presentation.save(output)
        encoded_decks.append(output.getvalue())

    target_size = max(map(len, encoded_decks))
    encoded_decks = [deck + b"\0" * (target_size - len(deck)) for deck in encoded_decks]
    assert len(encoded_decks[0]) == len(encoded_decks[1])
    path = tmp_path / "same-name.pptx"
    path.write_bytes(encoded_decks[0])
    first = read_pptx(path)
    path.write_bytes(encoded_decks[1])
    second = read_pptx(path)

    first_visual = next(element for element in first.elements if element.needs_vision)
    second_visual = next(element for element in second.elements if element.needs_vision)
    assert first_visual.element_id != second_visual.element_id


def test_plain_text_tracks_markdown_headings_and_character_spans(tmp_path):
    path = tmp_path / "notes.md"
    path.write_text("# 标题\n\n第一段内容。\n\n## 子节\n\n第二段内容。\n", encoding="utf-8")

    result = read_document(path)

    paragraphs = [e for e in result.elements if e.kind is ElementKind.PARAGRAPH]
    assert paragraphs[0].locator.section_path == ("标题",)
    assert paragraphs[1].locator.section_path == ("标题", "子节")
    assert paragraphs[0].locator.char_range is not None
    assert paragraphs[0].locator.page is None


def test_an_empty_source_is_diagnosed(tmp_path):
    path = tmp_path / "blank.txt"
    path.write_text("   \n\n", encoding="utf-8")

    result = read_document(path)

    assert result.elements == ()
    assert result.diagnostics[0].code is DiagnosticCode.EMPTY_SOURCE


def test_the_registry_covers_the_formats_the_old_reader_did(tmp_path):
    suffixes = set(supported_suffixes())
    assert {".pdf", ".docx", ".pptx", ".txt", ".md", ".csv"} <= suffixes


@pytest.mark.parametrize("suffix", (".doc", ".ppt"))
def test_legacy_office_requires_the_explicit_conversion_bridge(tmp_path, suffix):
    path = tmp_path / f"legacy{suffix}"
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")

    with pytest.raises(UnsupportedLegacyOfficeFormat) as exc:
        read_document(path)

    assert exc.value.reason_code == "unsupported_legacy_office"


def _tiny_png(tmp_path):
    image = pytest.importorskip("PIL.Image")
    path = tmp_path / "tiny.png"
    image.new("RGB", (8, 8), color=(255, 0, 0)).save(path)
    return path


def _move_shape_to(shape, target_shapes) -> None:
    """将不带关系的测试形状移入布局/母版形状树。"""

    cloned = copy.deepcopy(shape._element)
    cloned.xpath(".//p:cNvPr")[0].set("id", str(target_shapes._next_shape_id))
    target_shapes._spTree.insert_element_before(cloned, "p:extLst")
    shape._element.getparent().remove(shape._element)
