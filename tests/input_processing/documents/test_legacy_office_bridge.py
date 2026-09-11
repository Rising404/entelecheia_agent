from __future__ import annotations

import shutil
import subprocess
import os
from pathlib import Path

import pytest

from personagraph.input_processing.documents import (
    PreparedDocumentIngest,
    prepare_document_path,
)
from personagraph.input_processing.documents.contracts import (
    DiagnosticCode,
    ProcessingAdmissionStatus,
)
from personagraph.input_processing.documents.readers import read_document
from personagraph.input_processing.documents.readers.legacy_office import (
    LegacyOfficeBridge,
    MacOsSandboxExecLegacyOfficeRunner,
)


OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


class _CopyConvertedFixture:
    def __init__(self, converted_fixture: Path) -> None:
        self.converted_fixture = converted_fixture
        self.calls = 0

    def convert(
        self,
        *,
        soffice_path: Path,
        source_path: Path,
        output_dir: Path,
        profile_dir: Path,
        target_suffix: str,
        timeout_seconds: int,
    ) -> Path:
        self.calls += 1
        assert soffice_path.is_file()
        assert source_path.read_bytes().startswith(OLE_MAGIC)
        assert source_path.parent.name == "source"
        assert profile_dir.parent == output_dir.parent
        assert timeout_seconds == 60
        target = output_dir / f"{source_path.stem}{target_suffix}"
        shutil.copyfile(self.converted_fixture, target)
        return target


class _EscapingRunner(_CopyConvertedFixture):
    def convert(self, **kwargs) -> Path:  # type: ignore[no-untyped-def]
        outside = kwargs["output_dir"].parent / "escaped.docx"
        shutil.copyfile(self.converted_fixture, outside)
        return outside


class _TimeoutRunner(_CopyConvertedFixture):
    def convert(self, **kwargs) -> Path:  # type: ignore[no-untyped-def]
        raise subprocess.TimeoutExpired(cmd="soffice", timeout=60)


def _executable(tmp_path: Path) -> Path:
    executable = tmp_path / "soffice"
    executable.write_text("test converter identity", encoding="utf-8")
    executable.chmod(0o700)
    return executable


def _legacy_source(tmp_path: Path, suffix: str) -> Path:
    source = tmp_path / f"legacy{suffix}"
    source.write_bytes(OLE_MAGIC + b"binary-office-fixture")
    return source


def _docx_fixture(tmp_path: Path) -> Path:
    import docx

    path = tmp_path / "converted.docx"
    document = docx.Document()
    document.add_heading("Legacy Word", level=1)
    document.add_paragraph("The binary Word source was converted and read.")
    document.save(path)
    return path


def _pptx_fixture(tmp_path: Path) -> Path:
    from pptx import Presentation

    path = tmp_path / "converted.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Legacy PowerPoint"
    slide.placeholders[1].text = "The binary deck was converted and read."
    deck.save(path)
    return path


@pytest.mark.parametrize(
    ("suffix", "fixture_factory", "expected_text"),
    (
        (".doc", _docx_fixture, "binary Word source"),
        (".ppt", _pptx_fixture, "binary deck"),
    ),
)
def test_explicit_bridge_converts_then_uses_the_native_typed_reader(
    tmp_path: Path,
    suffix: str,
    fixture_factory,
    expected_text: str,
) -> None:
    runner = _CopyConvertedFixture(fixture_factory(tmp_path))
    bridge = LegacyOfficeBridge(
        soffice_path=_executable(tmp_path),
        runner=runner,
    )

    result = read_document(
        _legacy_source(tmp_path, suffix),
        legacy_office_bridge=bridge,
    )

    assert result.admission_status is ProcessingAdmissionStatus.PARTIAL
    assert expected_text in "\n".join(item.text or "" for item in result.elements)
    assert result.diagnostics[-1].code is DiagnosticCode.FALLBACK_READER_USED
    assert "sandboxed-libreoffice" in str(result.processor)
    assert runner.calls == 1


def test_prepare_service_admits_a_bridge_result_without_mutating_the_source(
    tmp_path: Path,
) -> None:
    source = _legacy_source(tmp_path, ".doc")
    original = source.read_bytes()
    bridge = LegacyOfficeBridge(
        soffice_path=_executable(tmp_path),
        runner=_CopyConvertedFixture(_docx_fixture(tmp_path)),
    )

    prepared = prepare_document_path(
        str(source),
        legacy_office_bridge=bridge,
    )

    assert isinstance(prepared, PreparedDocumentIngest)
    assert prepared.processing_status == "partial"
    assert prepared.mime == "doc"
    assert source.read_bytes() == original


def test_bridge_rejects_converter_output_outside_its_private_output_directory(
    tmp_path: Path,
) -> None:
    bridge = LegacyOfficeBridge(
        soffice_path=_executable(tmp_path),
        runner=_EscapingRunner(_docx_fixture(tmp_path)),
    )

    result = bridge.read(_legacy_source(tmp_path, ".doc"))

    assert result.elements == ()
    assert result.diagnostics[0].code is DiagnosticCode.CORRUPT_SOURCE


def test_bridge_projects_conversion_timeout_as_a_typed_hard_limit(
    tmp_path: Path,
) -> None:
    bridge = LegacyOfficeBridge(
        soffice_path=_executable(tmp_path),
        runner=_TimeoutRunner(_docx_fixture(tmp_path)),
    )

    result = bridge.read(_legacy_source(tmp_path, ".doc"))

    assert result.elements == ()
    assert result.diagnostics[0].code is DiagnosticCode.LIMIT_REACHED


def test_bridge_rejects_a_renamed_non_compound_file_before_process_dispatch(
    tmp_path: Path,
) -> None:
    runner = _CopyConvertedFixture(_docx_fixture(tmp_path))
    bridge = LegacyOfficeBridge(
        soffice_path=_executable(tmp_path),
        runner=runner,
    )
    source = tmp_path / "renamed.doc"
    source.write_text("not a binary Office document", encoding="utf-8")

    result = bridge.read(source)

    assert result.elements == ()
    assert result.diagnostics[0].code is DiagnosticCode.CORRUPT_SOURCE
    assert runner.calls == 0


@pytest.mark.skipif(
    not os.environ.get("PERSONAGRAPH_TEST_LEGACY_OFFICE_SOFFICE"),
    reason="real LibreOffice integration is explicitly opt-in",
)
@pytest.mark.parametrize(
    ("source_suffix", "target_suffix", "fixture_factory", "export_filter", "text"),
    (
        (".docx", ".doc", _docx_fixture, "doc:MS Word 97", "binary Word source"),
        (
            ".pptx",
            ".ppt",
            _pptx_fixture,
            "ppt:MS PowerPoint 97",
            "binary deck",
        ),
    ),
)
def test_real_libreoffice_binary_document_is_read_inside_the_macos_sandbox(
    tmp_path: Path,
    source_suffix: str,
    target_suffix: str,
    fixture_factory,
    export_filter: str,
    text: str,
) -> None:
    soffice = Path(os.environ["PERSONAGRAPH_TEST_LEGACY_OFFICE_SOFFICE"])
    modern = fixture_factory(tmp_path)
    assert modern.suffix == source_suffix
    legacy_dir = tmp_path / "legacy-output"
    setup_profile = tmp_path / "setup-profile"
    legacy_dir.mkdir()
    setup_profile.mkdir()
    created = subprocess.run(  # 可信夹具创建过程，不属于被测桥接器
        [
            str(soffice),
            "--headless",
            f"-env:UserInstallation={setup_profile.as_uri()}",
            "--convert-to",
            export_filter,
            "--outdir",
            str(legacy_dir),
            str(modern),
        ],
        check=False,
        stdin=subprocess.DEVNULL,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        timeout=60,
    )
    legacy = legacy_dir / f"converted{target_suffix}"
    assert created.returncode == 0
    assert legacy.read_bytes().startswith(OLE_MAGIC)
    bridge = LegacyOfficeBridge(
        soffice_path=soffice,
        runner=MacOsSandboxExecLegacyOfficeRunner(),
    )

    result = bridge.read(legacy)

    assert result.admission_status is ProcessingAdmissionStatus.PARTIAL
    assert text in "\n".join(
        item.text or "" for item in result.elements
    )
