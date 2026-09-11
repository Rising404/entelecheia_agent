"""面向模型、针对单个冻结工作区根目录的格式专属观察。"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import pytest
from jsonschema import Draft202012Validator
import personagraph.input_processing.vision.imaging.payload as visual_payload

from personagraph.workspace.storage.context import current
from personagraph.workspace.files import (
    FileSource,
    WorkspaceFileAuthority,
)
from personagraph.tools.execution import (
    ResolvedInvocation,
    ToolBusinessFailure,
    ToolExecutor,
)
from personagraph.input_processing.vision.providers import UnavailableVisionModelAdapter
from personagraph.input_processing.vision.imaging import (
    build_prepared_visual_artifact_receipt,
)
from personagraph.input_processing.vision.contracts import (
    VisionCapabilitySnapshot,
    VisionObservation,
    VisionPurpose,
    VisionResult,
    VisionStatus,
)
from personagraph.tools.documents.format_observation_tools import (
    build_external_visual_analysis_tool_registrations,
    build_format_observation_tool_registrations,
)
from personagraph.tools.workspace.workspace_tools import FrozenWorkspaceToolBoundary
from personagraph.session import store as session_store


EXPECTED_LOCAL_TOOL_IDS = (
    "read_text",
    "read_pdf_text",
    "read_word",
    "read_slides",
    "inspect_image",
)
EXPECTED_ANALYSIS_TOOL_IDS = (
    "analyze_image",
    "analyze_pdf_page",
)


def _plain(value):
    if hasattr(value, "items"):
        return {key: _plain(item) for key, item in value.items()}
    if isinstance(value, (tuple, list)):
        return [_plain(item) for item in value]
    return value


class _SemanticAdapter:
    def __init__(self, *, transmits: bool = False) -> None:
        self.transmits_externally = transmits
        self.seen = []
        self.payload_prefixes: list[bytes] = []
        self.payload_sha256: list[str] = []

    def capabilities(self) -> VisionCapabilitySnapshot:
        return VisionCapabilitySnapshot(
            available=True,
            provider="test",
            model="test-vision",
            endpoint_identity="test:vision",
            processor_fingerprint="test-vision@1",
            supported_purposes=tuple(VisionPurpose),
        )

    def analyze(self, request) -> VisionResult:
        self.seen.append(request)
        assert request.prepared_payload is not None
        payload = request.prepared_payload.data
        self.payload_prefixes.append(payload[:8])
        self.payload_sha256.append(hashlib.sha256(payload).hexdigest())
        return VisionResult(
            status=VisionStatus.COMPLETED,
            provider="test",
            model="test-vision",
            endpoint_identity="test:vision",
            processor_fingerprint="test-vision@1",
            input_sha256=request.image_sha256,
            observations=(
                VisionObservation(
                    observation_id=f"observation-{request.page}",
                    kind=request.purpose.value,
                    text=f"Semantic observation for page {request.page}.",
                    uncertainty=0.1,
                ),
            ),
        )


def _tools(root: Path, *, vision_adapter=None):
    boundary = FrozenWorkspaceToolBoundary(session_id="session-1", root=root)
    registrations = (
        *build_format_observation_tool_registrations(
            boundary,


        ),
        *build_external_visual_analysis_tool_registrations(
            boundary,
            vision_adapter=vision_adapter or UnavailableVisionModelAdapter(),

        ),
    )
    return {registration.tool_id: registration for registration in registrations}


def test_all_format_path_contracts_distinguish_relative_path_from_display_name(
    tmp_path: Path,
) -> None:
    tools = _tools(tmp_path)

    for tool_id in (*EXPECTED_LOCAL_TOOL_IDS, *EXPECTED_ANALYSIS_TOOL_IDS):
        spec = tools[tool_id].spec
        assert "工作区相对路径" in spec.description
        path_schema = spec.input_schema["properties"]["path"]
        assert "工作区相对路径" in path_schema["description"]
        assert "name" in path_schema["description"]


def _execute(registration, arguments):
    outcome = ToolExecutor().execute(
        ResolvedInvocation(registration=registration, arguments=arguments)
    )
    assert outcome.status.value == "succeeded", outcome.error
    return dict(outcome.result or {})


def _make_pdf(path: Path, pages: tuple[str, ...]) -> None:
    from reportlab.pdfgen.canvas import Canvas

    canvas = Canvas(str(path), pagesize=(300, 220))
    for text in pages:
        canvas.drawString(36, 160, text)
        canvas.showPage()
    canvas.save()


def _make_docx(path: Path) -> None:
    from docx import Document

    document = Document()
    document.add_heading("Word title", level=1)
    document.add_paragraph("Word body evidence")
    document.save(path)


def _make_docx_with_maximum_heading(path: Path) -> tuple[str, str]:
    from docx import Document

    heading = "H" * 20_000
    body_parts = (
        "first body evidence after the maximum heading",
        "second body evidence after the maximum heading",
        "third body evidence after the maximum heading",
    )
    document = Document()
    document.add_heading(heading, level=1)
    for body in body_parts:
        document.add_paragraph(body)
    document.save(path)
    return heading, "".join(body_parts)


def _make_docx_with_combined_metadata_pressure(
    path: Path,
    image_path: Path,
) -> str:
    from docx import Document

    document = Document()
    document.add_heading("H" * 20_000, level=1)
    body_parts = tuple("🧭" * 500 for _ in range(16))
    for body in body_parts:
        document.add_paragraph(body)
    for _ in range(16):
        document.add_picture(str(image_path))
    document.save(path)
    return "".join(body_parts)


def _make_pptx(path: Path) -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Slide title"
    slide.placeholders[1].text = "Slide body evidence"
    presentation.save(path)


def _make_png(path: Path) -> None:
    from PIL import Image

    Image.new("RGB", (40, 24), "white").save(path, format="PNG")


def _make_jpeg(path: Path) -> None:
    from PIL import Image

    Image.new("RGB", (40, 24), "white").save(path, format="JPEG")




def test_local_builder_returns_five_exact_readonly_registrations(tmp_path: Path):
    registrations = build_format_observation_tool_registrations(
        FrozenWorkspaceToolBoundary(session_id="session-1", root=tmp_path),

    )

    assert tuple(item.tool_id for item in registrations) == EXPECTED_LOCAL_TOOL_IDS
    assert {
        effect.action.value
        for registration in registrations
        for effect in registration.effect_profile.effects
    } == {"read"}
    for registration in registrations:
        Draft202012Validator.check_schema(_plain(registration.spec.input_schema))
        Draft202012Validator.check_schema(_plain(registration.spec.output_schema))
        assert registration.execution_profile.max_output_bytes <= 96 * 1024


def test_local_model_facing_tools_do_not_expose_pdf_rendering_probe(
    tmp_path: Path,
):
    assert "render_pdf_pages" not in _tools(tmp_path)


def test_external_analysis_is_declared_as_transmit_without_tainting_local_reads(
    tmp_path: Path,
):
    tools = _tools(tmp_path, vision_adapter=_SemanticAdapter(transmits=True))

    for tool_id in EXPECTED_ANALYSIS_TOOL_IDS:
        assert {
            effect.action.value for effect in tools[tool_id].effect_profile.effects
        } == {"read", "transmit", "update"}
    for tool_id in EXPECTED_LOCAL_TOOL_IDS:
        assert {
            effect.action.value for effect in tools[tool_id].effect_profile.effects
        } == {"read"}
        read_effect = tools[tool_id].effect_profile.effects[0]
        assert read_effect.scope_kind.value == "workspace"
        assert read_effect.default_scope == str(tmp_path.resolve())
    properties = _plain(tools["inspect_image"].spec.input_schema)["properties"]
    assert {"analyze", "purpose", "region"}.isdisjoint(properties)
    for tool_id in EXPECTED_ANALYSIS_TOOL_IDS:
        Draft202012Validator.check_schema(_plain(tools[tool_id].spec.input_schema))
        Draft202012Validator.check_schema(_plain(tools[tool_id].spec.output_schema))
        read_effect = next(
            effect
            for effect in tools[tool_id].effect_profile.effects
            if effect.action.value == "read"
        )
        assert read_effect.scope_kind.value == "workspace"
        assert read_effect.default_scope == str(tmp_path.resolve())
        transmit_effect = next(
            effect
            for effect in tools[tool_id].effect_profile.effects
            if effect.action.value == "transmit"
        )
        assert transmit_effect.scope_kind.value == "session"
        assert transmit_effect.default_scope == "session-1"


def test_read_text_has_a_lossless_continuation_cursor_and_explicit_coverage(
    tmp_path: Path,
):
    (tmp_path / "notes.md").write_text(
        "# Evidence\n\nalpha beta gamma\n\nsecond block",
        encoding="utf-8",
    )
    tool = _tools(tmp_path)["read_text"]

    first = _execute(
        tool,
        {
            "path": "notes.md",
            "max_chars": 5,
            "max_elements": 1,
        },
    )

    assert first["format"] == "text"
    assert first["elements"][0]["text"] == "Evide"
    assert first["truncated"] is True
    assert first["coverage"]["source_status"] == "complete"
    assert first["coverage"]["selection_complete"] is False
    assert first["next_cursor"] == {"element_offset": 0, "character_offset": 5}

    resumed = _execute(
        tool,
        {
            "path": "notes.md",
            "element_offset": 0,
            "character_offset": 5,
            "max_chars": 100,
            "max_elements": 10,
        },
    )
    assert resumed["elements"][0]["text"] == "nce"
    assert "alpha beta gamma" in {item["text"] for item in resumed["elements"]}


def test_default_cursor_can_resume_across_one_maximum_reader_element(tmp_path: Path):
    original = "x" * 20_000
    (tmp_path / "long.txt").write_text(original, encoding="utf-8")
    tool = _tools(tmp_path)["read_text"]

    observed = ""
    cursor = {}
    for _ in range(10):
        result = _execute(tool, {"path": "long.txt", **cursor})
        observed += "".join(item["text"] for item in result["elements"])
        cursor = dict(result["next_cursor"] or {})
        if not cursor:
            break

    assert observed == original
    assert cursor == {}


def test_word_maximum_heading_metadata_stays_bounded_across_continuation(
    tmp_path: Path,
):
    heading, body = _make_docx_with_maximum_heading(tmp_path / "long-heading.docx")
    registration = _tools(tmp_path)["read_word"]

    observed = ""
    cursor = {}
    results = []
    for _ in range(4):
        outcome = ToolExecutor().execute(
            ResolvedInvocation(
                registration=registration,
                arguments={
                    "path": "long-heading.docx",
                    "max_chars": 8_000,
                    "max_elements": 16,
                    **cursor,
                },
            )
        )
        assert outcome.status.value == "succeeded", outcome.error
        result = _plain(outcome.result or {})
        results.append(result)
        assert (
            len(
                json.dumps(
                    result,
                    ensure_ascii=False,
                    allow_nan=False,
                    separators=(",", ":"),
                ).encode("utf-8")
            )
            <= registration.execution_profile.max_output_bytes
        )
        observed += "".join(item["text"] for item in result["elements"])
        cursor = dict(result["next_cursor"] or {})
        if not cursor:
            break

    assert observed == heading + body
    assert results[0]["next_cursor"] == {
        "element_offset": 0,
        "character_offset": 8_000,
    }
    assert results[1]["next_cursor"] == {
        "element_offset": 0,
        "character_offset": 16_000,
    }
    assert any(
        diagnostic["code"] == "result_metadata_truncated"
        for result in results
        for diagnostic in result["diagnostics"]
    )
    assert all(result["coverage"]["selection_complete"] is False for result in results)


def test_word_combined_metadata_budget_returns_an_exact_continuation(
    tmp_path: Path,
):
    _make_png(tmp_path / "pixel.png")
    body = _make_docx_with_combined_metadata_pressure(
        tmp_path / "combined-pressure.docx",
        tmp_path / "pixel.png",
    )
    registration = _tools(tmp_path)["read_word"]

    outcome = ToolExecutor().execute(
        ResolvedInvocation(
            registration=registration,
            arguments={
                "path": "combined-pressure.docx",
                "element_offset": 1,
                "max_chars": 8_000,
                "max_elements": 16,
            },
        )
    )

    assert outcome.status.value == "succeeded", outcome.error
    first = _plain(outcome.result or {})
    assert any(
        item["code"] == "result_output_truncated" for item in first["diagnostics"]
    )
    assert first["truncated"] is True
    assert first["coverage"]["selection_complete"] is False
    assert first["next_cursor"] is not None

    observed = "".join(item["text"] for item in first["elements"])
    cursor = dict(first["next_cursor"])
    for _ in range(16):
        result = _execute(
            registration,
            {
                "path": "combined-pressure.docx",
                "max_chars": 8_000,
                "max_elements": 16,
                **cursor,
            },
        )
        observed += "".join(item["text"] for item in result["elements"])
        cursor = dict(result["next_cursor"] or {})
        if not cursor:
            break

    assert observed == body
    assert cursor == {}


def test_maximum_multibyte_text_page_stays_inside_the_workrun_result_budget(
    tmp_path: Path,
):
    original = "🧭" * 8_000
    (tmp_path / "unicode.txt").write_text(original, encoding="utf-8")

    result = _execute(
        _tools(tmp_path)["read_text"],
        {"path": "unicode.txt", "max_chars": 8_000, "max_elements": 1},
    )

    assert result["elements"][0]["text"] == original
    assert result["truncated"] is False


@pytest.mark.parametrize("escape", ("../outside.txt", "/etc/passwd"))
@pytest.mark.parametrize(
    "tool_id",
    (*EXPECTED_LOCAL_TOOL_IDS, *EXPECTED_ANALYSIS_TOOL_IDS),
)
def test_no_tool_can_widen_the_frozen_workspace_root(
    tmp_path: Path, escape: str, tool_id: str
):
    tool = _tools(tmp_path)[tool_id]

    with pytest.raises(ToolBusinessFailure) as raised:
        tool.handler({"path": escape})

    assert raised.value.error.code == "workspace_path_blocked"


def test_read_text_rejects_mixed_case_sensitive_directory_name(tmp_path: Path):
    hidden = tmp_path / ".GiT"
    hidden.mkdir()
    (hidden / "hidden.md").write_text("must stay hidden", encoding="utf-8")

    with pytest.raises(ToolBusinessFailure) as raised:
        _tools(tmp_path)["read_text"].handler({"path": ".GiT/hidden.md"})

    assert raised.value.error.code == "workspace_path_blocked"


def test_a_symlink_cannot_escape_the_frozen_workspace_root(tmp_path: Path):
    outside = tmp_path.parent / "format-observation-secret.txt"
    outside.write_text("secret", encoding="utf-8")
    link = tmp_path / "secret.txt"
    try:
        link.symlink_to(outside)
    except OSError:
        pytest.skip("symlinks are unavailable on this host")
    try:
        with pytest.raises(ToolBusinessFailure) as raised:
            _tools(tmp_path)["read_text"].handler({"path": "secret.txt"})
        assert raised.value.error.code == "workspace_path_blocked"
    finally:
        outside.unlink(missing_ok=True)


def test_pdf_text_can_be_narrowed_to_pages(
    tmp_path: Path,
):
    _make_pdf(
        tmp_path / "paper.pdf",
        (
            "page one has enough evidence",
            "page two has enough evidence",
            "page three has enough evidence",
        ),
    )
    tools = _tools(tmp_path)

    text = _execute(
        tools["read_pdf_text"],
        {
            "path": "paper.pdf",
            "start_page": 2,
            "end_page": 2,
            "max_chars": 8_000,
            "max_elements": 16,
        },
    )
    assert text["coverage"]["physical_pages"] == 3
    assert text["coverage"]["selected_pages"] == (2,)
    assert any("page two" in item["text"] for item in text["elements"])
    assert all("page one" not in item["text"] for item in text["elements"])


def test_word_and_slide_readers_preserve_text_and_slide_rendering_is_an_honest_gap(
    tmp_path: Path,
):
    _make_docx(tmp_path / "brief.docx")
    _make_pptx(tmp_path / "deck.pptx")
    tools = _tools(tmp_path)

    word = _execute(tools["read_word"], {"path": "brief.docx"})
    assert {item["text"] for item in word["elements"]} >= {
        "Word title",
        "Word body evidence",
    }

    slides = _execute(
        tools["read_slides"],
        {"path": "deck.pptx", "render_visuals": True},
    )
    assert any("Slide title" in item["text"] for item in slides["elements"])
    assert slides["rendering"] == {
        "requested": True,
        "status": "unavailable",
        "reason_code": "slide_render_backend_not_configured",
    }
    assert any(
        item["code"] == "slide_render_backend_unavailable"
        for item in slides["diagnostics"]
    )


def test_inspect_image_returns_only_pixel_metadata_and_keeps_visual_semantics_open(
    tmp_path: Path,
):
    _make_png(tmp_path / "figure.png")

    result = _execute(
        _tools(tmp_path)["inspect_image"],
        {"path": "figure.png", "detail": "low"},
    )

    assert result["format"] == "image"
    assert result["coverage"]["needs_vision"] is True
    assert result["images"][0]["width"] == 40
    assert result["images"][0]["height"] == 24
    assert "data_base64" not in result["images"][0]
    assert any(item["code"] == "page_needs_vision" for item in result["diagnostics"])


def test_local_visual_tools_never_invoke_an_external_adapter(tmp_path: Path):
    image_path = tmp_path / "figure.png"
    _make_png(image_path)
    adapter = _SemanticAdapter(transmits=True)
    tools = _tools(tmp_path, vision_adapter=adapter)

    inspected = _execute(
        tools["inspect_image"],
        {"path": "figure.png", "detail": "low"},
    )
    assert inspected["analysis"]["status"] == "not_requested"
    assert "data_base64" not in json.dumps(_plain(inspected), sort_keys=True)
    assert (
        "include_image_data"
        not in _plain(tools["inspect_image"].spec.input_schema)["properties"]
    )
    assert adapter.seen == []


@pytest.mark.parametrize("suffix", (".png", ".jpg"))
def test_analyze_image_invokes_exactly_one_adjustable_provider_read(
    tmp_path: Path,
    suffix: str,
):
    image_path = tmp_path / f"figure{suffix}"
    (_make_png if suffix == ".png" else _make_jpeg)(image_path)
    adapter = _SemanticAdapter(transmits=True)
    tool = _tools(
        tmp_path,
        vision_adapter=adapter,

    )["analyze_image"]

    analyzed = _execute(
        tool,
        {
            "path": image_path.name,
            "purpose": "chart",
            "detail": "low",
            "region": "page",
        },
    )
    assert analyzed["analysis"]["status"] == "completed"
    assert analyzed["status"] != "gap"
    assert analyzed["coverage"]["needs_vision"] is False
    assert analyzed["observations"][0]["observation"].startswith("Semantic observation")
    assert len(adapter.seen) == 1
    assert adapter.seen[0].purpose is VisionPurpose.CHART
    assert adapter.seen[0].detail.value == "low"
    assert adapter.seen[0].region.value == "page"


def test_analyze_image_requires_explicit_purpose_detail_and_region(tmp_path: Path):
    tool = _tools(tmp_path)["analyze_image"]
    schema = _plain(tool.spec.input_schema)

    errors = list(Draft202012Validator(schema).iter_errors({"path": "figure.png"}))
    assert errors
    assert list(
        Draft202012Validator(schema).iter_errors(
            {
                "path": "figure.png",
                "purpose": "general",
                "detail": "low",
                "region": "page",
                "include_image_data": True,
            }
        )
    )


def test_unavailable_vision_is_a_typed_semantic_gap(tmp_path: Path):
    _make_png(tmp_path / "figure.png")
    result = _execute(
        _tools(tmp_path, vision_adapter=UnavailableVisionModelAdapter())[
            "analyze_image"
        ],
        {
            "path": "figure.png",
            "purpose": "general",
            "detail": "standard",
            "region": "page",
        },
    )

    assert result["analysis"]["status"] == "unavailable"
    assert result["observations"][0]["failure_code"] == "vision_provider_unavailable"
    assert result["coverage"]["needs_vision"] is True


def test_analyze_pdf_page_invokes_exactly_one_selected_provider_read(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
):
    pdf_path = tmp_path / "paper.pdf"
    _make_pdf(
        pdf_path,
        ("page one has enough evidence", "page two has enough evidence"),
    )
    adapter = _SemanticAdapter(transmits=True)
    rendered_pages: list[int | None] = []
    original_render = visual_payload._render_pdf_region

    def counting_render(path, request, level):
        rendered_pages.append(request.locator.page)
        return original_render(path, request, level)

    monkeypatch.setattr(visual_payload, "_render_pdf_region", counting_render)
    result = _execute(
        _tools(
            tmp_path,
            vision_adapter=adapter,

        )["analyze_pdf_page"],
        {
            "path": "paper.pdf",
            "pages": [2],
            "purpose": "general",
            "detail": "low",
            "region": "page",
        },
    )

    assert result["analysis"]["status"] == "completed"
    assert tuple(item["page"] for item in result["observations"]) == (2,)
    assert tuple(request.page for request in adapter.seen) == (2,)
    assert rendered_pages == [2]
    assert result["coverage"]["needs_vision"] is False
    request = adapter.seen[0]
    source_sha256 = hashlib.sha256(pdf_path.read_bytes()).hexdigest()
    assert Path(request.image_path) == pdf_path.resolve()
    assert Path(request.image_path).suffix == ".pdf"
    assert request.source_sha256 == source_sha256
    assert request.image_sha256 == source_sha256
    assert request.prepared_payload is not None
    assert request.prepared_payload.source_sha256 == source_sha256
    assert request.pixel_size == request.prepared_payload.pixel_size
    assert adapter.payload_prefixes == [b"\x89PNG\r\n\x1a\n"]
    assert request.prepared_payload.sent_sha256 == adapter.payload_sha256[0]
    assert request.image_sha256 != adapter.payload_sha256[0]
    receipt = build_prepared_visual_artifact_receipt(
        request,
        request.prepared_payload,
    )
    recipe = json.loads(receipt.canonical_render_recipe)
    assert recipe["request"]["source_kind"] == "pdf-region"
    assert recipe["request"]["locator"]["page"] == 2


def test_registered_project_pdf_keeps_implicit_authority_on_original_source(
    tmp_path: Path,
    partitioned_project_state,
):
    root = tmp_path / "project"
    root.mkdir()
    session_id = session_store.create_session(
        "Entelecheia",
        working_dir=str(root),
    )
    pdf_path = root / "paper.pdf"
    _make_pdf(pdf_path, ("managed attachment page",))
    adapter = _SemanticAdapter(transmits=True)

    with session_store.session_database_scope(session_id):
        database = current()
        assert database is not None
        WorkspaceFileAuthority(database).register_path(
            "paper.pdf",
            source=FileSource.USER_UPLOAD,
            media_type="application/pdf",
        )
        result = _execute(
            _tools(
                root,
                vision_adapter=adapter,

            )["analyze_pdf_page"],
            {
                "path": "paper.pdf",
                "pages": [1],
                "purpose": "general",
                "detail": "low",
                "region": "page",
            },
        )

    assert result["analysis"]["status"] == "completed"
    assert len(adapter.seen) == 1
    request = adapter.seen[0]
    assert request.disclosure_receipt_id.startswith("auto_visual_egress_")
    assert Path(request.image_path) == pdf_path.resolve()
    assert Path(request.image_path).suffix == ".pdf"
    assert request.prepared_payload is not None
    assert request.prepared_payload.mime_type == "image/png"


def test_analyze_pdf_page_accepts_only_unique_explicit_page_batches(tmp_path: Path):
    _make_pdf(tmp_path / "paper.pdf", ("page one", "page two"))
    adapter = _SemanticAdapter(transmits=True)
    tool = _tools(tmp_path, vision_adapter=adapter)["analyze_pdf_page"]
    schema = _plain(tool.spec.input_schema)
    valid = {
        "path": "paper.pdf",
        "pages": [1, 2],
        "purpose": "general",
        "detail": "low",
        "region": "page",
    }

    assert not list(Draft202012Validator(schema).iter_errors(valid))
    for extra in (
        {"start_page": 1},
        {"end_page": 2},
        {"max_pages": 2},
        {"page": 1},
        {"pages": []},
        {"pages": [1, 1]},
        {"pages": [1, 2, 3, 4]},
        {"pages": [0]},
        {"pages": [True]},
    ):
        arguments = {**valid, **extra}
        assert list(Draft202012Validator(schema).iter_errors(arguments))
        outcome = ToolExecutor().execute(
            ResolvedInvocation(registration=tool, arguments=arguments)
        )
        assert outcome.status.value == "rejected"
    assert adapter.seen == []


def test_analyze_pdf_page_reads_selected_pages_once_in_requested_order(tmp_path: Path):
    _make_pdf(tmp_path / "paper.pdf", ("first page", "second page", "third page"))
    adapter = _SemanticAdapter(transmits=True)
    result = _execute(
        _tools(tmp_path, vision_adapter=adapter)["analyze_pdf_page"],
        {
            "path": "paper.pdf", "pages": [3, 1, 2], "purpose": "general",
            "detail": "low", "region": "page",
        },
    )
    assert [request.page for request in adapter.seen] == [3, 1, 2]
    assert tuple(result["coverage"]["selected_pages"]) == (3, 1, 2)
    assert [item["page"] for item in result["observations"]] == [3, 1, 2]
    assert result["analysis"]["requested_units"] == 3
    assert result["analysis"]["resolved_units"] == 3
    assert result["status"] == "completed"


@pytest.mark.parametrize("pages", [[1, 3], [2, 2], [], [1, 2, 3, 4]])
def test_analyze_pdf_page_prevalidates_whole_batch_before_provider(tmp_path: Path, pages):
    _make_pdf(tmp_path / "paper.pdf", ("first page", "second page"))
    adapter = _SemanticAdapter(transmits=True)
    tool = _tools(tmp_path, vision_adapter=adapter)["analyze_pdf_page"]
    with pytest.raises(ToolBusinessFailure):
        tool.handler({
            "path": "paper.pdf", "pages": pages, "purpose": "general",
            "detail": "low", "region": "page",
        })
    assert adapter.seen == []


def test_analyze_pdf_page_preserves_successes_around_provider_failure(tmp_path: Path):
    _make_pdf(tmp_path / "paper.pdf", ("first page", "second page", "third page"))

    class PartialAdapter(_SemanticAdapter):
        def analyze(self, request):
            result = super().analyze(request)
            if request.page != 2:
                return result
            return VisionResult(
                status=VisionStatus.FAILED, provider="test", model="test-vision",
                endpoint_identity="test:vision", processor_fingerprint="test-vision@1",
                input_sha256=request.image_sha256, failure_code="provider_unavailable",
                unresolved_gap_refs=(request.source_unit_id,),
            )

    adapter = PartialAdapter(transmits=True)
    result = _execute(
        _tools(tmp_path, vision_adapter=adapter)["analyze_pdf_page"],
        {
            "path": "paper.pdf", "pages": [1, 2, 3], "purpose": "general",
            "detail": "low", "region": "page",
        },
    )
    assert [item["status"] for item in result["observations"]] == [
        "completed", "failed", "completed",
    ]
    assert result["observations"][1]["failure_code"] == "provider_unavailable"
    assert result["status"] == "partial"
    assert result["analysis"]["resolved_units"] == 2
    assert result["analysis"]["requested_units"] == 3


def test_pdf_page_batch_has_explicit_contract_and_aggregate_time_budget(tmp_path: Path):
    tools = _tools(tmp_path)
    pdf_tool = tools["analyze_pdf_page"]
    assert "Attempt" in pdf_tool.spec.description
    assert "同一 Attempt 中多次调用" in pdf_tool.spec.description
    assert "按 calls 顺序串行执行" in pdf_tool.spec.description
    assert "最多调用本工具一次" not in pdf_tool.spec.description
    assert "最多包含一个受保护调用" not in pdf_tool.spec.description
    assert "1..3" in pdf_tool.spec.description
    assert pdf_tool.execution_profile.default_timeout_s == 360.0
    assert pdf_tool.execution_profile.hard_timeout_s == 360.0
    assert tools["analyze_image"].execution_profile.default_timeout_s == 120.0


def test_pdf_page_batch_stops_before_next_page_when_source_changes(tmp_path: Path):
    target = tmp_path / "paper.pdf"
    _make_pdf(target, ("first page", "second page"))

    class ChangingAdapter(_SemanticAdapter):
        def analyze(self, request):
            result = super().analyze(request)
            _make_pdf(target, ("replacement first page", "replacement second page"))
            return result

    adapter = ChangingAdapter(transmits=True)
    tool = _tools(tmp_path, vision_adapter=adapter)["analyze_pdf_page"]
    with pytest.raises(ToolBusinessFailure) as exc:
        tool.handler({
            "path": "paper.pdf", "pages": [1, 2], "purpose": "general",
            "detail": "low", "region": "page",
        })
    assert exc.value.error.code == "source_changed_during_read"
    assert [request.page for request in adapter.seen] == [1]


def test_remote_semantic_read_uses_default_egress_without_consent_state(tmp_path: Path):
    _make_png(tmp_path / "figure.png")
    adapter = _SemanticAdapter(transmits=True)
    result = _execute(
        _tools(
            tmp_path,
            vision_adapter=adapter,

        )["analyze_image"],
        {
            "path": "figure.png",
            "purpose": "general",
            "detail": "low",
            "region": "page",
        },
    )

    assert result["observations"][0]["status"] == "completed"
    assert len(adapter.seen) == 1
    assert adapter.seen[0].disclosure_receipt_id.startswith("auto_visual_egress_")


@pytest.mark.parametrize(
    ("tool_id", "suffix"),
    (("read_word", ".doc"), ("read_slides", ".ppt")),
)
def test_legacy_office_without_a_bridge_is_an_honest_gap(
    tmp_path: Path, tool_id: str, suffix: str
):
    (tmp_path / f"legacy{suffix}").write_bytes(b"legacy-office-placeholder")

    result = _execute(_tools(tmp_path)[tool_id], {"path": f"legacy{suffix}"})

    assert result["status"] == "gap"
    assert result["diagnostics"][0]["code"] == "legacy_office_bridge_unavailable"


def test_wrong_format_is_a_typed_non_retryable_request_error(tmp_path: Path):
    (tmp_path / "notes.txt").write_text("hello", encoding="utf-8")

    with pytest.raises(ToolBusinessFailure) as raised:
        _tools(tmp_path)["read_pdf_text"].handler({"path": "notes.txt"})

    assert raised.value.error.code == "format_mismatch"
    assert raised.value.error.details["actual_suffix"] == ".txt"
