"""DOCX 和 PPTX reader。

两种格式都具有旧契约无法表达的结构：Word heading 层级和幻灯片顺序。它们分别填写
``section_path`` 与 ``page``，并将几何信息留空；任一格式都无须了解另一格式的行为。

这些原生 reader 会主动报告无法解释的视觉单元。因此，嵌入图片或图表会成为已定位的
``IMAGE`` 元素并附带 ``PAGE_NEEDS_VISION`` 诊断，绝不会表现为看似完整的文本文档。
"""

from __future__ import annotations

import io
from collections.abc import Iterable
from pathlib import Path
from xml.etree import ElementTree

from ...files import (
    OoxmlFailureKind,
    OoxmlKind,
    OoxmlValidationError,
    read_validated_ooxml,
)
from ..contracts import (
    DiagnosticCode,
    DocumentElement,
    DocumentLocator,
    DocumentNonTextKind,
    DocumentNonTextUnit,
    DocumentPageInventoryStatus,
    DocumentPageManifest,
    DocumentPageRecord,
    DocumentPageState,
    ElementKind,
    ProcessingDiagnostic,
    ProcessingResult,
    ProcessorFingerprint,
    make_element_id,
)


DOCX_READER = ProcessorFingerprint("python-docx", "5")
PPTX_READER = ProcessorFingerprint("python-pptx", "7")

MAX_ELEMENT_CHARS = 20000
MAX_ELEMENTS = 20000


class _ElementEmitter:
    """在唯一元素创建边界应用两个全局 reader 限制。"""

    def __init__(self, source_key: str) -> None:
        self.source_key = source_key
        self.elements: list[DocumentElement] = []
        self.diagnostics: list[ProcessingDiagnostic] = []
        self.nontext_by_page: dict[int, list[DocumentNonTextUnit]] = {}
        self.budget_exhausted = False

    def emit_text(
        self,
        text: str,
        kind: ElementKind,
        *,
        ordinal: int,
        page: int | None = None,
        section_path: tuple[str, ...] = (),
    ) -> tuple[int, bool]:
        """对一个逻辑块生成无损且确定的分段。"""

        text = text.strip()
        if not text:
            return ordinal, True
        segmented = len(text) > MAX_ELEMENT_CHARS
        for start in range(0, len(text), MAX_ELEMENT_CHARS):
            end = min(start + MAX_ELEMENT_CHARS, len(text))
            stored_text = text[start:end]
            locator = DocumentLocator(
                page=page,
                ordinal=ordinal,
                section_path=section_path,
                char_range=(start, end) if segmented else None,
            )
            if not self._reserve(locator):
                return ordinal, False
            self.elements.append(DocumentElement(
                element_id=make_element_id(self.source_key, locator, stored_text),
                kind=kind,
                text=stored_text,
                locator=locator,
            ))
            ordinal += 1
        return ordinal, True

    def emit_visual(
        self,
        *,
        ordinal: int,
        page: int | None = None,
        section_path: tuple[str, ...] = (),
        detail: str,
        nontext_kind: DocumentNonTextKind = DocumentNonTextKind.FIGURE,
        text_element_ids: tuple[str, ...] = (),
    ) -> tuple[int, bool]:
        locator = DocumentLocator(
            page=page,
            ordinal=ordinal,
            section_path=section_path,
        )
        if not self._reserve(locator):
            return ordinal, False
        self.elements.append(DocumentElement(
            element_id=make_element_id(self.source_key, locator, None),
            kind=ElementKind.IMAGE,
            text=None,
            locator=locator,
            needs_vision=True,
        ))
        self.diagnostics.append(ProcessingDiagnostic(
            DiagnosticCode.PAGE_NEEDS_VISION,
            locator,
            detail=detail,
        ))
        if page is not None:
            self.record_nontext(
                page=page,
                kind=nontext_kind,
                locator=locator,
                element_id=self.elements[-1].element_id,
                text_element_ids=text_element_ids,
                requires_visual_read=True,
            )
        return ordinal + 1, True

    def record_nontext(
        self,
        *,
        page: int,
        kind: DocumentNonTextKind,
        locator: DocumentLocator,
        element_id: str | None,
        text_element_ids: tuple[str, ...] = (),
        requires_visual_read: bool,
    ) -> None:
        """将一个稳定非文本单元绑定到其精确幻灯片证据。"""

        identity = make_element_id(
            self.source_key,
            locator,
            f"<pptx-nontext:{kind.value}>",
        )
        self.nontext_by_page.setdefault(page, []).append(DocumentNonTextUnit(
            unit_id=f"unit_{identity.removeprefix('el_')}",
            kind=kind,
            source_pages=(page,),
            text_element_ids=text_element_ids,
            element_id=element_id,
            locator=locator,
            requires_visual_read=requires_visual_read,
        ))

    def _reserve(self, locator: DocumentLocator) -> bool:
        if len(self.elements) < MAX_ELEMENTS:
            return True
        if not self.budget_exhausted:
            self.diagnostics.append(ProcessingDiagnostic(
                DiagnosticCode.LIMIT_REACHED,
                locator,
                detail="element budget exhausted",
            ))
            self.budget_exhausted = True
        return False


def read_docx(path: Path) -> ProcessingResult:
    """按源顺序读取 Word 块，并保留 heading 层级。"""

    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        source = read_validated_ooxml(path, expected_kind=OoxmlKind.DOCX)
    except OoxmlValidationError as exc:
        return _ooxml_failure_result(DOCX_READER, exc)
    except PermissionError as exc:
        return _office_read_failure(DOCX_READER, DiagnosticCode.PERMISSION_DENIED, exc)
    except OSError as exc:
        return _office_read_failure(DOCX_READER, DiagnosticCode.CORRUPT_SOURCE, exc)
    try:
        document = docx.Document(io.BytesIO(source.payload))
    except Exception as exc:
        return _office_read_failure(DOCX_READER, DiagnosticCode.CORRUPT_SOURCE, exc)
    emitter = _ElementEmitter(source.source_sha256)
    section: list[str] = []
    ordinal = 0

    for label, story in _distinct_docx_stories(document, family="header"):
        ordinal, complete = _emit_docx_story(
            story,
            emitter,
            ordinal=ordinal,
            section_path=(label,),
        )
        if not complete:
            break
    else:
        complete = True

    for block in _iter_docx_blocks(document) if complete else ():
        if isinstance(block, Paragraph):
            text = _docx_paragraph_text(block).strip()
            depth = _heading_depth(block) if text else 0
            if depth:
                section = [*section[: depth - 1], text]
                kind = ElementKind.HEADING
            else:
                kind = ElementKind.PARAGRAPH
            ordinal, complete = emitter.emit_text(
                text,
                kind,
                ordinal=ordinal,
                section_path=tuple(section),
            )
        elif isinstance(block, Table):
            ordinal, complete = emitter.emit_text(
                _render_docx_table(block),
                ElementKind.TABLE,
                ordinal=ordinal,
                section_path=tuple(section),
            )
        else:  # pragma: no cover - 依赖契约只会产生这两种值
            continue
        if not complete:
            break

        for detail in _docx_visual_details(block):
            ordinal, complete = emitter.emit_visual(
                ordinal=ordinal,
                section_path=tuple(section),
                detail=detail,
            )
            if not complete:
                break
        if not complete:
            break

    if complete:
        for alt_chunk_index, _alt_chunk in enumerate(
            _docx_alt_chunks(document),
            start=1,
        ):
            ordinal, complete = emitter.emit_visual(
                ordinal=ordinal,
                section_path=(f"altChunk:{alt_chunk_index}",),
                detail=(
                    "embedded Word alternative-format content was not "
                    "interpreted"
                ),
            )
            if not complete:
                break

    if complete:
        note_stories, missing_note_refs = _docx_note_stories(document)
        for label, blocks in note_stories:
            ordinal, complete = _emit_docx_blocks(
                blocks,
                emitter,
                ordinal=ordinal,
                section_path=(label,),
            )
            if not complete:
                break
        if complete and missing_note_refs:
            emitter.diagnostics.append(ProcessingDiagnostic(
                DiagnosticCode.CORRUPT_SOURCE,
                detail=(
                    "referenced DOCX note/comment bodies are missing: "
                    + ",".join(missing_note_refs)
                ),
            ))

    if complete:
        for label, story in _distinct_docx_stories(document, family="footer"):
            ordinal, complete = _emit_docx_story(
                story,
                emitter,
                ordinal=ordinal,
                section_path=(label,),
            )
            if not complete:
                break

    if not emitter.elements:
        emitter.diagnostics.append(ProcessingDiagnostic(DiagnosticCode.EMPTY_SOURCE))
    return ProcessingResult(
        tuple(emitter.elements), DOCX_READER, tuple(emitter.diagnostics)
    )


def read_pptx(path: Path) -> ProcessingResult:
    """读取演示文稿；每张幻灯片都是一页，视觉形状保持为显式 gap。"""

    from pptx import Presentation

    try:
        source = read_validated_ooxml(path, expected_kind=OoxmlKind.PPTX)
    except OoxmlValidationError as exc:
        return _ooxml_failure_result(PPTX_READER, exc)
    except PermissionError as exc:
        return _office_read_failure(PPTX_READER, DiagnosticCode.PERMISSION_DENIED, exc)
    except OSError as exc:
        return _office_read_failure(PPTX_READER, DiagnosticCode.CORRUPT_SOURCE, exc)
    try:
        presentation = Presentation(io.BytesIO(source.payload))
    except Exception as exc:
        return _office_read_failure(PPTX_READER, DiagnosticCode.CORRUPT_SOURCE, exc)
    emitter = _ElementEmitter(source.source_sha256)
    physical_slide_count = len(presentation.slides)

    for slide_index, slide in enumerate(presentation.slides, start=1):
        ordinal = 0
        complete = True
        background_role = _pptx_effective_image_background_role(slide)
        if background_role is not None:
            ordinal, complete = emitter.emit_visual(
                page=slide_index,
                ordinal=ordinal,
                section_path=(background_role,),
                detail="PowerPoint image background requires visual interpretation",
                nontext_kind=DocumentNonTextKind.FIGURE,
            )
        if not complete:
            break
        for source_role, shape in _iter_pptx_inherited_shapes(slide):
            ordinal, complete = _emit_pptx_shape(
                shape,
                emitter,
                page=slide_index,
                ordinal=ordinal,
                report_visual=True,
                section_path=(source_role,),
            )
            if not complete:
                break
        if not complete:
            break
        for shape in slide.shapes:
            ordinal, complete = _emit_pptx_shape(
                shape,
                emitter,
                page=slide_index,
                ordinal=ordinal,
                report_visual=True,
            )
            if not complete:
                break
        if not complete:
            break

        if getattr(slide, "has_notes_slide", False):
            notes_slide = slide.notes_slide
            notes_frame = getattr(notes_slide, "notes_text_frame", None)
            notes = str(getattr(notes_frame, "text", "") or "").strip()
            ordinal, complete = emitter.emit_text(
                notes,
                ElementKind.PARAGRAPH,
                page=slide_index,
                ordinal=ordinal,
                section_path=("speaker_notes",),
            )
            if not complete:
                break
            for shape in _pptx_additional_notes_shapes(notes_slide):
                ordinal, complete = _emit_pptx_shape(
                    shape,
                    emitter,
                    page=slide_index,
                    ordinal=ordinal,
                    report_visual=True,
                    section_path=("speaker_notes_extra",),
                )
                if not complete:
                    break
            if not complete:
                break

        ordinal, complete = _emit_pptx_review_comments(
            slide,
            emitter,
            page=slide_index,
            ordinal=ordinal,
        )
        if not complete:
            break

    page_manifest = _build_pptx_page_manifest(
        physical_slide_count=physical_slide_count,
        emitter=emitter,
    )
    if not emitter.elements:
        emitter.diagnostics.append(ProcessingDiagnostic(DiagnosticCode.EMPTY_SOURCE))
    return ProcessingResult(
        tuple(emitter.elements),
        PPTX_READER,
        tuple(emitter.diagnostics),
        page_manifest=page_manifest,
    )


def _ooxml_failure_result(
    processor: ProcessorFingerprint,
    error: OoxmlValidationError,
) -> ProcessingResult:
    code = {
        OoxmlFailureKind.CORRUPT: DiagnosticCode.CORRUPT_SOURCE,
        OoxmlFailureKind.ENCRYPTED: DiagnosticCode.PASSWORD_REQUIRED,
        OoxmlFailureKind.LIMIT: DiagnosticCode.LIMIT_REACHED,
    }[error.kind]
    return ProcessingResult(
        elements=(),
        processor=processor,
        diagnostics=(ProcessingDiagnostic(code, detail=error.detail),),
    )


def _office_read_failure(
    processor: ProcessorFingerprint,
    code: DiagnosticCode,
    error: BaseException,
) -> ProcessingResult:
    return ProcessingResult(
        elements=(),
        processor=processor,
        diagnostics=(ProcessingDiagnostic(code, detail=type(error).__name__),),
    )


def _iter_docx_blocks(document) -> Iterable[object]:
    """按顺序生成段落/表格，包括结构化控件。

    python-docx 会刻意从 ``paragraphs`` 和 ``iter_inner_content`` 中省略 ``w:sdt`` 内容
    控件。这些控件常见于表单和模板，因此遍历小型块级 OOXML 词汇是避免静默丢失其文本的
    唯一方式。
    """

    element = getattr(document, "element", None)
    container = getattr(element, "body", None)
    if container is None:
        container = getattr(document, "_element")

    yield from _walk_docx_blocks(container, owner=document)


def _walk_docx_blocks(container, *, owner) -> Iterable[object]:
    """遍历当前可见块包装器，但不深入段落。"""

    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in container.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, owner)
        elif isinstance(child, CT_Tbl):
            yield Table(child, owner)
        else:
            local_name = str(child.tag).rsplit("}", 1)[-1]
            if local_name in {"del", "moveFrom"}:
        # Word 当前可见视图排除已删除/移出修订内容。插入/移入/content-control 及其他块
        # 包装器会在下方递归解包。
                continue
            yield from _walk_docx_blocks(child, owner=owner)


def _docx_paragraph_text(paragraph) -> str:
    """渲染 Word 当前可见 run 树，同时不丢失包装器。

    ``Paragraph.text`` 会省略嵌套在修订和内容控件包装器中的 run。这些包装器常见于经评审的
    合同和表单文档，因此本函数直接遍历小型可见文本词汇。已删除和移出内容被刻意排除。
    """

    from docx.oxml.ns import qn

    text_tag = qn("w:t")
    tab_tag = qn("w:tab")
    break_tags = {qn("w:br"), qn("w:cr")}
    excluded_tags = {qn("w:del"), qn("w:moveFrom")}

    def visit(node) -> Iterable[str]:
        if node.tag in excluded_tags:
            return
        if node.tag == text_tag:
            if node.text:
                yield str(node.text)
            return
        if node.tag == tab_tag:
            yield "\t"
            return
        if node.tag in break_tags:
            yield "\n"
            return
        for child in node.iterchildren():
            yield from visit(child)

    return "".join(visit(paragraph._element))


def _docx_alt_chunks(document) -> tuple[object, ...]:
    """为替代格式导入建立清单，但不获取或执行它们。"""

    from docx.oxml.ns import qn

    return tuple(document.element.body.iter(qn("w:altChunk")))


def _docx_note_stories(document):
    """将被引用的脚注、尾注和评审评论解析为 story。"""

    from docx.oxml import parse_xml
    from docx.oxml.ns import qn

    specs = {
        "footnote": ("footnotes", qn("w:footnoteReference"), qn("w:footnote")),
        "endnote": ("endnotes", qn("w:endnoteReference"), qn("w:endnote")),
        "comment": ("comments", qn("w:commentReference"), qn("w:comment")),
    }
    stories: list[tuple[str, tuple[object, ...]]] = []
    missing: list[str] = []
    body = document.element.body
    for label, (relationship_tail, reference_tag, item_tag) in specs.items():
        referenced = {
            str(node.get(qn("w:id")))
            for node in body.iter(reference_tag)
            if node.get(qn("w:id")) is not None
        }
        if not referenced:
            continue
        part = next((
            rel.target_part
            for rel in document.part.rels.values()
            if str(rel.reltype).endswith(f"/{relationship_tail}")
        ), None)
        if part is None:
            missing.extend(f"{label}:{item_id}" for item_id in sorted(referenced))
            continue
        try:
            root = parse_xml(part.blob)
        except Exception:
            missing.extend(f"{label}:{item_id}" for item_id in sorted(referenced))
            continue
        found: set[str] = set()
        for item in root.iterchildren(item_tag):
            item_id = item.get(qn("w:id"))
            if item_id not in referenced:
                continue
            found.add(str(item_id))
            stories.append((
                f"{label}:{item_id}",
                tuple(_walk_docx_blocks(item, owner=document)),
            ))
        missing.extend(
            f"{label}:{item_id}" for item_id in sorted(referenced - found)
        )
    return tuple(stories), tuple(missing)


def _distinct_docx_stories(document, *, family: str):
    """即使 section 链接到页眉/页脚变体，也只生成每个变体一次。"""

    if family not in {"header", "footer"}:
        raise ValueError("DOCX story family must be header or footer")
    seen: set[tuple[str, str]] = set()
    for section in document.sections:
        variants = [("default", family)]
        if bool(getattr(section, "different_first_page_header_footer", False)):
            variants.append(("first", f"first_page_{family}"))
        if bool(getattr(
            document.settings,
            "odd_and_even_pages_header_footer",
            False,
        )):
            variants.append(("even", f"even_page_{family}"))
        for role, attribute in variants:
            story = getattr(section, attribute)
            part_name = str(getattr(story.part, "partname", ""))
            key = (role, part_name or str(id(story._element)))
            if key in seen:
                continue
            seen.add(key)
            yield f"{family}:{role}", story


def _emit_docx_story(
    story,
    emitter: _ElementEmitter,
    *,
    ordinal: int,
    section_path: tuple[str, ...],
) -> tuple[int, bool]:
    return _emit_docx_blocks(
        _iter_docx_blocks(story),
        emitter,
        ordinal=ordinal,
        section_path=section_path,
    )


def _emit_docx_blocks(
    blocks: Iterable[object],
    emitter: _ElementEmitter,
    *,
    ordinal: int,
    section_path: tuple[str, ...],
) -> tuple[int, bool]:
    from docx.table import Table

    for block in blocks:
        if isinstance(block, Table):
            text = _render_docx_table(block)
            kind = ElementKind.TABLE
        else:
            text = _docx_paragraph_text(block)
            kind = ElementKind.PARAGRAPH
        ordinal, complete = emitter.emit_text(
            text,
            kind,
            ordinal=ordinal,
            section_path=section_path,
        )
        if not complete:
            return ordinal, False
        for detail in _docx_visual_details(block):
            ordinal, complete = emitter.emit_visual(
                ordinal=ordinal,
                section_path=section_path,
                detail=detail,
            )
            if not complete:
                return ordinal, False
    return ordinal, True


def _docx_visual_details(block) -> tuple[str, ...]:
    """恰好描述一次顶层绘图、对象和 OMML 公式。"""

    from docx.oxml.ns import qn

    detail_by_tag = {
        qn("w:drawing"): "embedded Word drawing requires visual interpretation",
        qn("w:pict"): "embedded Word picture requires visual interpretation",
        qn("w:object"): "embedded Word object requires visual interpretation",
        qn("m:oMath"): "embedded Word formula requires visual interpretation",
        qn("m:oMathPara"): "embedded Word formula requires visual interpretation",
    }
    visual_tags = set(detail_by_tag)
    root = block._element
    details: list[str] = []
    for node in root.iter():
        if node.tag not in visual_tags:
            continue
        parent = node.getparent()
        nested = False
        while parent is not None and parent is not root:
            if parent.tag in visual_tags:
                nested = True
                break
            parent = parent.getparent()
        if not nested:
            details.append(detail_by_tag[node.tag])
    return tuple(details)


def _emit_pptx_shape(
    shape,
    emitter: _ElementEmitter,
    *,
    page: int,
    ordinal: int,
    report_visual: bool,
    section_path: tuple[str, ...] = (),
) -> tuple[int, bool]:
    """按稳定顺序生成一个形状的可读部分与未读视觉部分。"""

    shape_text_ids: tuple[str, ...] = ()
    if getattr(shape, "has_table", False):
        element_start = len(emitter.elements)
        ordinal, complete = emitter.emit_text(
            _render_pptx_table(shape.table),
            ElementKind.TABLE,
            page=page,
            ordinal=ordinal,
            section_path=section_path,
        )
        table_elements = tuple(emitter.elements[element_start:])
        if complete and table_elements:
            emitter.record_nontext(
                page=page,
                kind=DocumentNonTextKind.TABLE,
                locator=table_elements[0].locator,
                element_id=table_elements[0].element_id,
                text_element_ids=tuple(
                    element.element_id for element in table_elements
                ),
                requires_visual_read=False,
            )
        elif complete:
            ordinal, complete = emitter.emit_visual(
                page=page,
                ordinal=ordinal,
                section_path=section_path,
                detail="PowerPoint empty table requires visual interpretation",
                nontext_kind=DocumentNonTextKind.TABLE,
            )
    elif getattr(shape, "has_text_frame", False):
        element_start = len(emitter.elements)
        ordinal, complete = emitter.emit_text(
            str(getattr(shape, "text", "") or ""),
            ElementKind.PARAGRAPH,
            page=page,
            ordinal=ordinal,
            section_path=section_path,
        )
        shape_text_ids = tuple(
            element.element_id for element in emitter.elements[element_start:]
        )
    else:
        complete = True
    if not complete:
        return ordinal, False

    if report_visual and _pptx_shape_requires_vision(shape):
        name = _pptx_shape_type_name(shape).lower().replace("_", " ")
        ordinal, complete = emitter.emit_visual(
            page=page,
            ordinal=ordinal,
            section_path=section_path,
            detail=f"PowerPoint {name} requires visual interpretation",
            nontext_kind=_pptx_nontext_kind(shape),
            text_element_ids=shape_text_ids,
        )
        if not complete:
            return ordinal, False

    if _pptx_shape_type_name(shape) == "GROUP":
        # 一个复合 gap 核算子布局/视觉关系；子文本和表格仍有用，且不会重复 group gap。
        for child in shape.shapes:
            ordinal, complete = _emit_pptx_shape(
                child,
                emitter,
                page=page,
                ordinal=ordinal,
                report_visual=False,
                section_path=section_path,
            )
            if not complete:
                return ordinal, False
    return ordinal, True


def _pptx_shape_requires_vision(shape) -> bool:
    name = _pptx_shape_type_name(shape)
    if getattr(shape, "has_table", False):
        # 原生表格以类型化 TABLE 单元无损表示。
        return False
        # 插入图片占位符的图片在 python-pptx 中仍为 shape_type PLACEHOLDER；实际 OOXML tag
        # 才是可靠信号。
    element_tag = str(getattr(getattr(shape, "_element", None), "tag", ""))
    if element_tag.rsplit("}", 1)[-1] in {"pic", "oleObj"}:
        return True
    if name in {
        "CHART",
        "DIAGRAM",
        "EMBEDDED_OLE_OBJECT",
        "GROUP",
        "IGX_GRAPHIC",
        "LINKED_OLE_OBJECT",
        "LINKED_PICTURE",
        "MEDIA",
        "OLE_CONTROL_OBJECT",
        "PICTURE",
        "WEB_VIDEO",
    }:
        return True
    if name in {
        "AUTO_SHAPE",
        "CALLOUT",
        "CANVAS",
        "FREEFORM",
        "INK",
        "LINE",
        "TEXT_EFFECT",
    }:
        # 文本仍有用，但它不编码箭头方向、连接线、包含关系或其他视觉关系。
        return True
    if name in {"PLACEHOLDER", "TEXT_BOX"} and getattr(
        shape, "has_text_frame", False
    ):
        return False
    # 新增或不常见的 Office 形状类型并不能证明内容缺失。当 python-pptx 增加新 enum，或遇到
    # 此适配器尚不理解的控件/评论形状时，保守保留视觉 gap 比静默宣称幻灯片完整更安全。
    return True


def _iter_pptx_inherited_shapes(slide):
    """为一张幻灯片生成有效的非占位 master/layout 形状。"""

    if not _pptx_follows_master_graphics(slide):
        return
    layout = slide.slide_layout
    if _pptx_follows_master_graphics(layout):
        for shape in layout.slide_master.shapes:
            if not getattr(shape, "is_placeholder", False):
                yield "slide_master", shape
    for shape in layout.shapes:
        if not getattr(shape, "is_placeholder", False):
            yield "slide_layout", shape


def _pptx_additional_notes_shapes(notes_slide):
    """生成内置 notes 正文之外由用户创作的 notes 形状。"""

    for shape in notes_slide.shapes:
        if getattr(shape, "is_placeholder", False):
        # 内置幻灯片图像、notes 正文、日期/页脚和幻灯片编号占位符已被表示，或仅是演示装饰。
            continue
        yield shape


def _emit_pptx_review_comments(
    slide,
    emitter: _ElementEmitter,
    *,
    page: int,
    ordinal: int,
) -> tuple[int, bool]:
    """生成经典评审评论，并对未知评论 XML 如实失败。"""

    comment_relationships = tuple(
        relationship
        for relationship in slide.part.rels.values()
        if "comments" in str(getattr(relationship, "reltype", "")).lower()
    )
    for relationship_index, relationship in enumerate(comment_relationships, start=1):
        try:
            if bool(getattr(relationship, "is_external", False)):
                raise ValueError("external comment relationship")
            payload = bytes(relationship.target_part.blob)
            root = ElementTree.fromstring(payload)
            comments = tuple(
                node
                for node in root.iter()
                if str(node.tag).rsplit("}", 1)[-1] in {"cm", "comment"}
            )
            emitted = 0
            for comment_index, comment in enumerate(comments, start=1):
                text = "".join(
                    child.text or ""
                    for child in comment.iter()
                    if str(child.tag).rsplit("}", 1)[-1] in {"text", "t"}
                ).strip()
                if not text:
                    continue
                ordinal, complete = emitter.emit_text(
                    text,
                    ElementKind.PARAGRAPH,
                    page=page,
                    ordinal=ordinal,
                    section_path=(
                        "review_comment",
                        f"relationship_{relationship_index}",
                        f"comment_{comment_index}",
                    ),
                )
                emitted += 1
                if not complete:
                    return ordinal, False
            if comments and emitted:
                continue
        except Exception:
            pass
        emitter.diagnostics.append(ProcessingDiagnostic(
            DiagnosticCode.PARSER_PARTIAL,
            DocumentLocator(page=page),
            detail="PowerPoint review comments could not be safely extracted",
        ))
    return ordinal, True


def _pptx_effective_image_background_role(slide) -> str | None:
    """返回携带栅格背景的有效继承层。"""

    def explicit_background(owner):
        element = getattr(owner, "_element", None)
        common_slide_data = getattr(element, "cSld", None)
        if common_slide_data is None:
            return None
        return next(
            (
                child
                for child in common_slide_data.iterchildren()
                if str(child.tag).rsplit("}", 1)[-1] == "bg"
            ),
            None,
        )

    def contains_image(background) -> bool:
        return background is not None and any(
            str(node.tag).rsplit("}", 1)[-1] in {"blip", "blipFill"}
            for node in background.iter()
        )

    slide_background = explicit_background(slide)
    if slide_background is not None:
        return "slide_background" if contains_image(slide_background) else None
    if not bool(getattr(slide, "follow_master_background", True)):
        return None
    layout = slide.slide_layout
    layout_background = explicit_background(layout)
    if layout_background is not None:
        return "slide_layout_background" if contains_image(layout_background) else None
    master_background = explicit_background(layout.slide_master)
    if contains_image(master_background):
        return "slide_master_background"
    return None


def _pptx_follows_master_graphics(owner) -> bool:
    """存在时读取公开继承标志，否则读取 OOXML showMasterSp。"""

    for attribute in ("follow_master_graphics", "show_master_shapes"):
        try:
            value = getattr(owner, attribute)
        except (AttributeError, ValueError):
            continue
        if isinstance(value, bool):
            return value
    element = getattr(owner, "_element", None)
    raw = None if element is None else element.get("showMasterSp")
    if raw is None:
        return True
    normalized = str(raw).strip().lower()
    if normalized in {"1", "true", "on", "yes"}:
        return True
    if normalized in {"0", "false", "off", "no"}:
        return False
    # 未知词法值不足以构成公开继承内容的 authority。
    return False


def _pptx_nontext_kind(shape) -> DocumentNonTextKind:
    name = _pptx_shape_type_name(shape)
    if name in {
        "CHART",
        "EMBEDDED_OLE_OBJECT",
        "LINKED_OLE_OBJECT",
        "LINKED_PICTURE",
        "MEDIA",
        "OLE_CONTROL_OBJECT",
        "PICTURE",
        "WEB_VIDEO",
    }:
        return DocumentNonTextKind.FIGURE
    return DocumentNonTextKind.VECTOR_GRAPHICS


def _build_pptx_page_manifest(
    *,
    physical_slide_count: int,
    emitter: _ElementEmitter,
) -> DocumentPageManifest:
    """为完整物理幻灯片全集建立清单，包括空幻灯片。"""

    if physical_slide_count == 0:
        return DocumentPageManifest(
            physical_page_count=0,
            inventory_status=DocumentPageInventoryStatus.UNAVAILABLE,
            detector_fingerprint=f"{PPTX_READER}:physical-slides-v1",
            detector_capabilities=(),
            pages=(),
        )

    limit_pages = tuple(
        diagnostic.locator.page
        for diagnostic in emitter.diagnostics
        if diagnostic.code is DiagnosticCode.LIMIT_REACHED
        and diagnostic.locator.page is not None
    )
    if emitter.budget_exhausted and limit_pages:
        first_unreadable = min(limit_pages)
        for page_number in range(first_unreadable + 1, physical_slide_count + 1):
            emitter.diagnostics.append(ProcessingDiagnostic(
                DiagnosticCode.LIMIT_REACHED,
                DocumentLocator(page=page_number),
                detail="slide not inspected after element budget exhaustion",
            ))

    text_ids_by_page: dict[int, list[str]] = {
        page: [] for page in range(1, physical_slide_count + 1)
    }
    for element in emitter.elements:
        if (element.text or "").strip() and element.locator.page is not None:
            text_ids_by_page[element.locator.page].append(element.element_id)

    diagnostics_by_page: dict[int, list[ProcessingDiagnostic]] = {
        page: [] for page in range(1, physical_slide_count + 1)
    }
    for diagnostic in emitter.diagnostics:
        if diagnostic.locator.page in diagnostics_by_page:
            diagnostics_by_page[int(diagnostic.locator.page)].append(diagnostic)

    pages: list[DocumentPageRecord] = []
    for page_number in range(1, physical_slide_count + 1):
        text_ids = tuple(text_ids_by_page[page_number])
        nontext_units = tuple(emitter.nontext_by_page.get(page_number, ()))
        page_diagnostics = diagnostics_by_page[page_number]
        fatal = any(
            diagnostic.code is DiagnosticCode.LIMIT_REACHED
            for diagnostic in page_diagnostics
        )
        if not text_ids and not nontext_units and not fatal:
            empty = ProcessingDiagnostic(
                DiagnosticCode.PAGE_EMPTY,
                DocumentLocator(page=page_number),
                detail="slide has no extractable text or visual unit",
            )
            emitter.diagnostics.append(empty)
            page_diagnostics.append(empty)
        if fatal:
            state = DocumentPageState.UNREADABLE
        elif text_ids and nontext_units:
            state = DocumentPageState.MIXED
        elif text_ids:
            state = DocumentPageState.TEXT
        elif nontext_units:
            state = DocumentPageState.VISUAL_ONLY
        else:
            state = DocumentPageState.NO_EXTRACTABLE_CONTENT
        pages.append(DocumentPageRecord(
            page_number=page_number,
            state=state,
            text_element_ids=text_ids,
            nontext_units=nontext_units,
            diagnostics=tuple(page_diagnostics),
        ))

    return DocumentPageManifest(
        physical_page_count=physical_slide_count,
        inventory_status=(
            DocumentPageInventoryStatus.PARTIAL
            if emitter.budget_exhausted
            else DocumentPageInventoryStatus.COMPLETE
        ),
        detector_fingerprint=f"{PPTX_READER}:physical-slides-v1",
        detector_capabilities=tuple(sorted({
            "figure_inventory",
            "inherited_shape_inventory",
            "nontext_unit_inventory",
            "physical_page_inventory",
            "slide_shape_inventory",
            "table_inventory",
            "text_element_source_pages",
            "typed_page_diagnostics",
            "vector_graphics_inventory",
        })),
        pages=tuple(pages),
    )


def _pptx_shape_type_name(shape) -> str:
    shape_type = getattr(shape, "shape_type", None)
    name = getattr(shape_type, "name", None)
    if isinstance(name, str):
        return name
    return str(shape_type or "UNKNOWN").split(" ", 1)[0]


def _heading_depth(paragraph) -> int:
    style = getattr(getattr(paragraph, "style", None), "name", "") or ""
    if not style.lower().startswith("heading"):
        return 0
    tail = style.split()[-1]
    return int(tail) if tail.isdigit() else 1


def _render_docx_table(table) -> str:
    from docx.table import Table

    rows = []
    for row in table.rows:
        cells = []
        for cell in row.cells:
            parts: list[str] = []
            for block in _iter_docx_blocks(cell):
                if isinstance(block, Table):
                    nested = _render_docx_table(block)
                    if nested.strip():
                        parts.append(f"[nested table: {nested.replace(chr(10), ' / ')}]")
                else:
                    text = _docx_paragraph_text(block).strip()
                    if text:
                        parts.append(text)
            cells.append(" ".join(parts))
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)


def _render_pptx_table(table) -> str:
    rows = []
    for row in table.rows:
        cells = [(cell.text or "").replace("\n", " ").strip() for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows)
